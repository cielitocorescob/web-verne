"""Credenciales VERNE 2026 — v3: editorial redesign on the original master (Poppins, VERNE palette)."""
import copy, zipfile, shutil
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC, OUT, A = 'Credenciales_VERNE_2026.pptx', 'Credenciales_VERNE_2026_v3.pptx', 'assets/'
MEDIA = '/tmp/claude-0/-home-user-web-verne/7d4b3b84-7c70-578c-b33c-1b2bd060621f/scratchpad/u/ppt/media/'
p = Presentation(SRC)
S = p.slides

PURPLE, DARK, LAV, LAV2, LAVBG, GRAY, TXT, WHITE, MUTED, NAVY = (
    '6D4ABE', '25283D', 'C0B0F4', 'D8D4EA', 'F3F0FB', '6B6E80', '333333', 'FFFFFF', '9A9CAB', '171828')
MID = '8F72D6'


# ------------------------------------------------------------------ helpers
def I(v): return Emu(int(round(v * 914400)))
def inch(e): return e / 914400


def sh(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)


def rm(slide, *ids):
    for sid in ids:
        e = sh(slide, sid)._element
        e.getparent().remove(e)


def keep_only(slide, keep):
    for s in list(slide.shapes):
        if s.shape_id not in keep:
            s._element.getparent().remove(s._element)


def mv(s, x=None, y=None, w=None, h=None):
    if x is not None: s.left = I(x)
    if y is not None: s.top = I(y)
    if w is not None: s.width = I(w)
    if h is not None: s.height = I(h)
    return s


def to_back(slide, s):
    tree = slide.shapes._spTree
    tree.remove(s._element)
    tree.insert(2, s._element)


def _alpha(el, pct):
    clr = el.find('.//' + qn('a:srgbClr'))
    a = clr.makeelement(qn('a:alpha'), {'val': str(int(pct * 1000))})
    clr.append(a)


def text(slide, x, y, w, h, paras, align='l', anchor='t', after=0, line=None, bullet=None):
    """paras: list of paragraphs, each a list of runs (text, pt, color, bold[, spc])."""
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    for i, runs in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if after: par.space_after = Pt(after)
        if line: par.line_spacing = line
        if bullet:
            pPr = par._p.get_or_add_pPr()
            pPr.set('marL', str(int(0.2 * 914400))); pPr.set('indent', str(int(-0.2 * 914400)))
            bc = pPr.makeelement(qn('a:buClr'), {}); c = bc.makeelement(qn('a:srgbClr'), {'val': bullet}); bc.append(c)
            pPr.append(bc)
            pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
            pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•'}))
        for run in runs:
            t, pt, color, bold = run[:4]
            r = par.add_run(); r.text = t
            r.font.size = Pt(pt); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)
            rPr = r._r.get_or_add_rPr()
            if len(run) > 4: rPr.set('spc', str(run[4]))
            for tag in ('a:latin', 'a:ea', 'a:cs'):
                rPr.append(rPr.makeelement(qn(tag), {'typeface': 'Poppins'}))
    return tb


def T(slide, x, y, w, h, s, pt, color, bold=False, **kw):
    return text(slide, x, y, w, h, [[(s, pt, color, bold)]], **kw)


def kicker(slide, x, y, w, s, color=PURPLE):
    return text(slide, x, y, w, 0.28, [[(s.upper(), 10.5, color, True, 300)]])


def box(slide, x, y, w, h, fill, alpha=None, line=None, line_alpha=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    if alpha is not None: _alpha(s._element.spPr.find(qn('a:solidFill')), alpha)
    if line:
        s.line.color.rgb = RGBColor.from_string(line); s.line.width = Pt(0.75)
        if line_alpha is not None: _alpha(s._element.spPr.find(qn('a:ln')), line_alpha)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None: s.adjustments[0] = radius
    return s


def hline(slide, x, y, w, color, alpha=None, weight=0.012):
    return box(slide, x, y, w, weight, color, alpha)


def vline(slide, x, y, h, color, alpha=None):
    return box(slide, x, y, 0.012, h, color, alpha)


def chip(slide, x, y, w, h, label, fill, color, pt=10.5, alpha=None):
    box(slide, x, y, w, h, fill, alpha, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    T(slide, x, y, w, h, label, pt, color, True, align='c', anchor='m')


def pic(slide, path, x, y, w=None, h=None):
    kw = {}
    if w is not None: kw['width'] = I(w)
    if h is not None: kw['height'] = I(h)
    return slide.shapes.add_picture(path, I(x), I(y), **kw)


def logo_grid(slide, ids, x, y, cols, cw, ch, maxw, maxh, align='c'):
    """Lay existing logo pictures on a tidy grid, keeping aspect ratio."""
    for i, sid in enumerate(ids):
        s = sh(slide, sid)
        ar = s.width / s.height
        w, h = maxh * ar, maxh
        if w > maxw: w, h = maxw, maxw / ar
        cx, cy = x + (i % cols) * cw, y + (i // cols) * ch
        ox = (cw - w) / 2 if align == 'c' else 0
        mv(s, cx + ox, cy + (ch - h) / 2, w, h)


def dark_chrome(slide, n, logo=True):
    """Master logo/number/url are hidden under full-bleed images: redraw them in white."""
    if logo: pic(slide, A + 'logo_white.png', 11.97, 0.42, 1.05, 0.37)
    T(slide, 0.24, 7.06, 0.5, 0.25, str(n), 9, WHITE, True, align='r')
    T(slide, 11.02, 7.02, 2.0, 0.22, 'www.verne.la', 9, LAV2, align='r')


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


# Slide-level page-number boxes duplicate the master's ‹Nº›; drop them everywhere (we redraw where needed).
for n, sid in [(3, 49), (4, 24), (5, 34), (6, 26), (7, 18), (8, 21), (9, 34), (10, 37), (11, 37), (12, 39), (13, 26)]:
    rm(S[n - 1], sid)

# ================================================================== 1 · PORTADA (dark, photo)
s = S[0]
keep_only(s, {2, 6, 7})
pic(s, A + 'logo_white.png', 0.65, 0.55, 1.45, 0.51)
kicker(s, 0.65, 2.2, 6, 'Credenciales 2026', LAV)
t = sh(s, 6); mv(t, 0.65, 2.6, 6.6, 2.3)
for r in t.text_frame.paragraphs[0].runs: r.font.size = Pt(36)
hline(s, 0.65, 6.62, 0.6, LAV)
mv(sh(s, 7), 0.65, 6.78)
notes(s, """
Apertura. VERNE: Consultora Estratégica de Comunicación Digital, potenciada con Inteligencia Artificial.
www.verne.la
""")

# ================================================================== 2 · QUIÉNES SOMOS (light, big data)
s = S[1]
rm(s, 35)
t = sh(s, 4); mv(t, y=1.28, w=12, h=0.85)
for r in t.text_frame.paragraphs[0].runs: r.font.size = Pt(40)
mv(sh(s, 5), y=2.18)
for c, cn, lab, stat, desc, src, ln, act in [(6, 7, 8, 9, 10, 11, 12, 13), (15, 16, 17, 18, 19, 20, 21, 22), (24, 25, 26, 27, 28, 29, 30, 31)]:
    for i in (c, cn, lab): mv(sh(s, i), y=2.75)
    st = sh(s, stat); mv(st, y=3.18, h=1.2)
    for r in st.text_frame.paragraphs[0].runs: r.font.size = Pt(84)
    mv(sh(s, desc), y=4.45); mv(sh(s, src), y=5.15)
    mv(sh(s, ln), y=5.46); mv(sh(s, act), y=5.56, h=0.62)
    for r in sh(s, act).text_frame.paragraphs[0].runs: r.font.color.rgb = RGBColor.from_string(PURPLE)
for v in (14, 23): mv(sh(s, v), y=2.75, h=3.4)
mv(sh(s, 32), y=6.45); mv(sh(s, 33), y=6.45)
notes(s, """
La comunicación estratégica ya no vive solo en los medios. Hoy se juega en tres frentes:

1. Reputación y liderazgo — 3x más confianza cuando la información la comparte un directivo o colaborador que la cuenta institucional (Fuente: Edelman Trust Barometer). Posicionamos a la marca corporativa, a sus líderes y embajadores en el ecosistema digital.
2. Visibilidad y presencia digital — 70% de los consumidores peruanos usa IA para informarse antes de comprar o contratar (Fuente: Impronta Research, ago 2025). Logramos que te encuentren, citen y recomienden en buscadores, IA y medios.
3. Adopción y capacidades de IA — 90% de las áreas de Comunicación y Marketing en el mundo ya experimenta con IA (Fuente: McKinsey, 2026). Habilitamos a los equipos para usar IA generativa en su día a día.

Cierre: Ni solo tecnología ni solo comunicación tradicional.
""")

# ================================================================== 3 · CLIENTES (light, editorial logo wall)
s = S[2]
sectors = [('Energía, Minería e Hidrocarburos', list(range(9, 22))),
           ('Banca y Servicios Financieros', list(range(25, 30))),
           ('Consumo, Retail e Industria', list(range(33, 37))),
           ('Educación, Servicios e Infraestructura', list(range(40, 49)))]
keep_only(s, {sid for _, ids in sectors for sid in ids})
kicker(s, 0.65, 1.0, 4, 'Con quiénes trabajamos')
T(s, 0.65, 1.3, 3.8, 1.75, 'Confían en nosotros', 44, DARK, True, line=0.95)
T(s, 0.65, 3.2, 3.4, 0.9, 'Empresas líderes de sectores regulados y competitivos.', 15, GRAY)
X0, WR = 4.75, 7.93
y = 1.15
cw = WR / 9
for name, ids in sectors:
    hline(s, X0, y, WR, 'E3E1EA')
    T(s, X0, y + 0.14, WR, 0.3, name, 11, PURPLE, True)
    rows = (len(ids) + 8) // 9
    logo_grid(s, ids, X0, y + 0.5, 9, cw, 0.5, 0.74, 0.32)
    y += 0.5 + rows * 0.5 + 0.28
notes(s, """
Confían en nosotros: empresas líderes de sectores regulados y competitivos.
Sectores: Energía, Minería e Hidrocarburos · Banca y Servicios Financieros · Consumo, Retail e Industria · Educación, Servicios e Infraestructura.
""")

# ================================================================== 4 · SERVICIOS (dark, service map)
s = S[3]
keep_only(s, set())
bg = pic(s, MEDIA + 'image36.jpg', 0, 0, 13.333, 7.5); to_back(s, bg)
dark_chrome(s, 4)
kicker(s, 0.65, 1.0, 5, 'Qué hacemos', LAV)
T(s, 0.65, 1.28, 9, 0.8, 'Nuestros servicios', 40, WHITE, True)
cols = [(3.25, '1', 'Reputación y liderazgo'), (6.40, '2', 'Visibilidad y presencia digital'), (9.55, '3', 'Adopción y capacidades de IA')]
CW = 3.0
for x, n, name in cols:
    text(s, x, 2.3, CW, 0.7, [[(n, 30, LAV, True)], [(name, 13, WHITE, True)]], line=0.95)
box(s, 0.0, 4.92, 13.333, 1.75, PURPLE, alpha=16)
hline(s, 0.65, 3.12, 12.03, WHITE, alpha=22)
hline(s, 0.65, 4.92, 12.03, WHITE, alpha=22)
for y, lab, sub in [(3.32, 'Consultoría', '(lo hacemos por ti)'), (5.12, 'Formación', '(habilitamos a tu equipo)')]:
    text(s, 0.65, y, 2.4, 0.8, [[(lab.upper(), 12, WHITE, True, 300)], [(sub, 11, LAV2, False)]])
consult = [['Marca Directiva y Corporativa', 'Embajadores digitales'],
           ['Posicionamiento GEO', 'Marketing de Contenidos', 'Asuntos Públicos'],
           ['Diagnóstico de madurez digital', 'Agentic AI Studio', 'Agentic AI Reputation', 'GEO Rank', 'Data Pulse']]
form = [['Talleres de Marca Directiva y Embajadores'],
        ['Formación para periodistas', 'Political Intelligence AI'],
        ['Talleres de IA, data y automatización', 'Taller de IA para comunicación y MKT', 'Claude AI para empresas', 'Taller IA Agéntica']]
for (x, _, _), c, f in zip(cols, consult, form):
    text(s, x, 3.32, CW, 1.5, [[(i, 11.5, WHITE, False)] for i in c], after=3)
    text(s, x, 5.12, CW, 1.5, [[(i, 11.5, WHITE, False)] for i in f], after=3)
notes(s, """
Dos modalidades para cada uno de los tres frentes:
- CONSULTORÍA (lo hacemos por ti)
- FORMACIÓN (habilitamos a tu equipo)

1. Reputación y liderazgo — Consultoría: Marca Directiva y Corporativa; Embajadores digitales. Formación: Talleres de Marca Directiva y Embajadores.
2. Visibilidad y presencia digital — Consultoría: Posicionamiento GEO; Marketing de Contenidos; Asuntos Públicos. Formación: Formación para periodistas; Political Intelligence AI.
3. Adopción y capacidades de IA — Consultoría: Diagnóstico de madurez digital; Agentic AI Studio; Agentic AI Reputation; GEO Rank; Data Pulse. Formación: Talleres de IA, data y automatización; Taller de IA para comunicación y MKT; Claude AI para empresas; Taller IA Agéntica.
""")

# ================================================================== 5 · MARCA CORPORATIVA Y DIRECTIVA (photo right + caption)
s = S[4]
logos5 = list(range(23, 34))
keep_only(s, set(logos5))
pic(s, A + 'p5.jpg', 8.133, 0, 5.2, 7.5)
pic(s, A + 'fade_bottom.png', 8.133, 0, 5.2, 7.5)
pic(s, A + 'logo_white.png', 11.97, 0.42, 1.05, 0.37)
T(s, 8.6, 5.55, 4.3, 1.2, 'Los líderes cambian, la marca permanece.', 22, WHITE, True, anchor='b', line=1.0)
kicker(s, 0.65, 1.0, 7, 'Reputación y liderazgo · Consultoría')
text(s, 0.65, 1.28, 7.0, 1.3, [[('Marca Corporativa', 34, DARK, True)], [('y Marca Directiva', 34, DARK, True)]], line=0.95)
T(s, 0.65, 2.68, 6.9, 0.65, 'Construimos la marca de tu empresa a través de la voz de sus líderes.', 15, GRAY)
for x, tag, name, defi in [(0.65, 'Es lo que permanece', 'Marca Corporativa', 'La reputación y credibilidad de la empresa en el ecosistema digital.'),
                           (4.25, 'Es lo que la amplifica', 'Marca Directiva', 'Tus C-level, directivos y voceros dan rostro y voz a esa narrativa.')]:
    kicker(s, x, 3.6, 3.3, tag)
    T(s, x, 3.9, 3.3, 0.42, name, 18, DARK, True)
    T(s, x, 4.36, 3.2, 0.7, defi, 11.5, TXT)
vline(s, 4.0, 3.62, 1.4, 'E3E1EA')
hline(s, 0.65, 5.3, 7.0, 'E3E1EA')
x = 0.65
for i, (n, name) in enumerate([('01', 'Diagnóstico'), ('02', 'Estrategia'), ('03', 'Gestión')]):
    text(s, x, 5.45, 1.9, 0.4, [[(n + '  ', 14, PURPLE, True), (name, 14, DARK, True)]], anchor='m')
    if i < 2: T(s, x + 1.9, 5.45, 0.4, 0.4, '→', 14, MUTED, align='c', anchor='m')
    x += 2.35
kicker(s, 0.65, 6.17, 1.3, 'Clientes')
logo_grid(s, logos5, 1.85, 5.98, 6, 0.97, 0.42, 0.78, 0.26)
notes(s, """
Construimos la marca de tu empresa a través de la voz de sus líderes. Los líderes cambian, la marca permanece.

Marca Corporativa: la reputación y credibilidad de la empresa en el ecosistema digital. Es lo que permanece.
Marca Directiva: tus C-level, directivos y voceros dan rostro y voz a esa narrativa. Es lo que la amplifica.

Cómo trabajamos:
01 Diagnóstico — Analizamos la presencia digital de la empresa y de sus directivos, y la comparamos con la de sus pares en la industria.
02 Estrategia — Definimos la narrativa corporativa, los ejes de contenido y el rol de cada vocero, para que el mensaje sea de la empresa y no dependa de una sola persona.
03 Gestión — Creamos contenidos, acompañamos a los voceros y medimos el desempeño de la marca corporativa y de sus líderes.
""")

# ================================================================== 6 · EMBAJADORES (photo left, vertical method)
s = S[5]
logos6 = list(range(20, 26))
keep_only(s, set(logos6))
pic(s, A + 'p6.jpg', 0, 0, 5.2, 7.5)
T(s, 5.4, 7.06, 0.5, 0.25, '6', 9, TXT, True)
X = 5.85
kicker(s, X, 1.0, 6.8, 'Reputación y liderazgo · Consultoría')
T(s, X, 1.28, 6.8, 0.8, 'Embajadores Digitales', 40, DARK, True)
T(s, X, 2.2, 6.6, 0.95, 'Amplificamos la comunicación corporativa a través de tus colaboradores, quienes hablan como la voz más genuina de la marca.', 15, GRAY)
steps = [('01', 'Estrategia', 'Seleccionamos y activamos embajadores alineados a los objetivos de la marca.'),
         ('02', 'Acompañamiento', 'Buenas prácticas, benchmark y co-creación de contenidos, con soporte continuo.'),
         ('03', 'Medición', 'Data pública, matriz de evaluación e informe mensual de desempeño, con Data Pulse.')]
y = 3.35
for n, name, d in steps:
    hline(s, X, y, 6.83, 'E3E1EA')
    T(s, X, y + 0.14, 1.0, 0.62, n, 30, LAV, True)
    T(s, X + 1.15, y + 0.14, 5.6, 0.36, name, 15, DARK, True)
    T(s, X + 1.15, y + 0.48, 5.6, 0.4, d, 11.5, GRAY)
    y += 0.9
hline(s, X, y, 6.83, 'E3E1EA')
kicker(s, X, 6.38, 1.4, 'Clientes')
logo_grid(s, logos6, X + 1.35, 6.25, 6, 0.92, 0.42, 0.78, 0.26)
notes(s, """
Amplificamos la comunicación corporativa a través de tus colaboradores, quienes hablan como la voz más genuina de la marca.

01 Estrategia — Seleccionamos y activamos embajadores alineados a los objetivos de la marca.
02 Acompañamiento — Buenas prácticas, benchmark y co-creación de contenidos, con soporte continuo.
03 Medición — Data pública, matriz de evaluación e informe mensual de desempeño, con Data Pulse.
""")

# ================================================================== 7 · GEO (dark, AI-answer visual + timeline)
s = S[6]
keep_only(s, set())
bg = pic(s, MEDIA + 'image50.jpg', 0, 0, 13.333, 7.5); to_back(s, bg)
dark_chrome(s, 7)
kicker(s, 0.65, 1.0, 7, 'Visibilidad y presencia digital · Consultoría', LAV)
T(s, 0.65, 1.28, 7, 0.8, 'Posicionamiento GEO', 40, WHITE, True)
T(s, 0.65, 2.2, 6.4, 0.95, 'Logramos que la IA cite y recomiende a tu marca.', 22, WHITE, True, line=1.0)
T(s, 0.65, 3.2, 6.2, 0.62, 'Trabajamos para que ChatGPT, Google AI Overview, Perplexity y Claude te mencionen en tu categoría.', 12.5, LAV2)
text(s, 0.65, 4.0, 6.2, 0.36, [[('CLIENTES   ', 9.5, LAV, True, 300), ('Gloria · ISA', 12, WHITE, True)]])
# illustrative AI-answer card: the brand surfaces inside the engines' response
PX, PY, PW, PH = 7.75, 1.25, 4.93, 2.95
box(s, PX, PY, PW, PH, WHITE, alpha=6, line=WHITE, line_alpha=18, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
x = PX + 0.3
for name, w in [('ChatGPT', 0.86), ('Google AI Overview', 1.45), ('Perplexity', 0.95), ('Claude', 0.72)]:
    chip(s, x, PY + 0.3, w, 0.32, name, PURPLE, WHITE, 8.5, alpha=45)
    x += w + 0.08
for dy, w in [(0.95, 4.1), (1.2, 3.4)]:
    box(s, PX + 0.3, PY + dy, w, 0.09, WHITE, alpha=16, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
box(s, PX + 0.3, PY + 1.53, 1.1, 0.09, WHITE, alpha=16, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
chip(s, PX + 1.5, PY + 1.38, 1.15, 0.38, 'Tu marca', PURPLE, WHITE, 11)
box(s, PX + 2.75, PY + 1.53, 1.65, 0.09, WHITE, alpha=16, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
for dy, w in [(1.95, 3.8), (2.2, 2.6)]:
    box(s, PX + 0.3, PY + dy, w, 0.09, WHITE, alpha=16, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
T(s, PX + 0.3, PY + 2.45, PW - 0.6, 0.3, 'Visibilidad · posición · citaciones · share of voice', 9, LAV, True)
# proportional timeline (15 / 15 / 60 días)
hline(s, 0.65, 4.72, 12.03, WHITE, alpha=18)
T(s, 0.65, 4.86, 12.03, 0.35, 'Resultados medibles en visibilidad, posición, citaciones y share of voice en IA.', 12.5, WHITE, True)
gap, unit = 0.06, (12.03 - 0.12) / 90
x = 0.65
for n, name, days, fill in [('01', 'Diagnóstico', 15, LAV), ('02', 'Estrategia', 15, MID), ('03', 'Implementación', 60, PURPLE)]:
    w = days * unit
    box(s, x, 5.42, w, 0.3, fill)
    text(s, x, 5.85, w, 0.32, [[(n + '  ', 12, LAV, True), (name, 12, WHITE, True)]])
    text(s, x, 6.18, w, 0.55, [[(str(days), 24, WHITE, True), (' días', 11, LAV2, False)]])
    x += w + gap
notes(s, """
Logramos que la IA cite y recomiende a tu marca. Trabajamos para que ChatGPT, Google AI Overview, Perplexity y Claude te mencionen en tu categoría.
(La tarjeta de la derecha es una ilustración del objetivo: que tu marca aparezca en la respuesta de los motores de IA.)

01 Diagnóstico (15 días) — Una radiografía de cómo aparece tu marca frente a la competencia, con brechas y oportunidades.
02 Estrategia (15 días) — Definimos narrativas, arquitectura de contenidos y el recorrido del usuario en la IA.
03 Implementación (60 días) — Parrilla GEO, producción de contenidos y verificación de citación en las IAs.

Resultados medibles en visibilidad, posición, citaciones y share of voice en IA.
Clientes: Gloria · ISA.
""")

# ================================================================== 8 · DIAGNÓSTICO DE MADUREZ (light, maturity staircase)
s = S[7]
keep_only(s, set())
kicker(s, 0.65, 1.0, 6, 'Adopción y capacidades de IA · Consultoría')
T(s, 0.65, 1.28, 5.8, 1.3, 'Diagnóstico de Madurez Digital', 36, DARK, True, line=0.95)
T(s, 0.65, 2.72, 5.4, 0.9, 'Te decimos qué tan bien usan tus equipos la IA y dónde invertir primero.', 16, GRAY)
text(s, 0.65, 3.78, 5.6, 0.3, [[('CLIENTES   ', 9.5, PURPLE, True, 300), ('BCP / Credicorp · Primax', 12, DARK, True)]])
BX, BW, BG, BASE = 7.25, 0.92, 0.205, 4.2
for i, c in enumerate(['E3E1EA', LAV2, LAV, MID, PURPLE]):
    h = 0.5 + i * 0.38
    box(s, BX + i * (BW + BG), BASE - h, BW, h, c)
    T(s, BX + i * (BW + BG), BASE - 0.42, BW, 0.32, str(i + 1), 13, DARK if i < 3 else WHITE, True, align='c')
text(s, 7.25, 1.1, 3.0, 1.2, [[('5 ', 30, PURPLE, True), ('niveles', 14, DARK, True)],
                              [('7 ', 30, PURPLE, True), ('indicadores por área', 14, DARK, True)]], line=0.9)
T(s, 7.25, 4.3, 5.43, 0.25, 'Los niveles se inspiran en el modelo CMMI.', 9.5, GRAY, align='r')
hline(s, 0.65, 4.62, 12.03, 'E3E1EA')
steps = [('01', 'Medimos la adopción', 'A través de una encuesta digital representativa y entrevistas a profundidad con quienes deciden, habilitan y construyen soluciones.'),
         ('02', 'Comparamos', 'Ubicamos a cada área en una escala de 5 niveles según su uso de IA generativa, automatización y datos.'),
         ('03', 'Priorizamos', 'Dashboard interactivo, informe ejecutivo y segmentación de gerencias.')]
for i, (n, name, d) in enumerate(steps):
    x = 0.65 + i * 4.08
    text(s, x, 4.82, 3.8, 0.4, [[(n + '  ', 15, PURPLE, True), (name, 15, DARK, True)]])
    T(s, x, 5.28, 3.7, 0.95, d, 11.5, GRAY)
x = 0.65 + 2 * 4.08
for lbl, fill, col in [('0-3', LAV, DARK), ('3-9', MID, WHITE), ('9-18', PURPLE, WHITE)]:
    chip(s, x, 5.9, 0.82, 0.34, lbl, fill, col, 10.5)
    x += 0.88
T(s, x + 0.04, 5.9, 1.2, 0.34, 'meses', 10.5, GRAY, anchor='m')
notes(s, """
Te decimos qué tan bien usan tus equipos la IA y dónde invertir primero. Los niveles se inspiran en el modelo CMMI.

01 Medimos la adopción — A través de una encuesta digital representativa y entrevistas a profundidad con quienes deciden, habilitan y construyen soluciones.
02 Comparamos — Ubicamos a cada área en una escala de 5 niveles según su uso de IA generativa, automatización y datos. Cada una se mide con 7 indicadores, como autonomía, impacto y capacitación.
03 Priorizamos — Dashboard interactivo, informe ejecutivo, segmentación de gerencias y hoja de ruta a 0-3, 3-9 y 9-18 meses.

Clientes: BCP / Credicorp · Primax.
""")

# ================================================================== 9 · CUATRO SOLUCIONES (dark, product showcase)
s = S[8]
rm(s, 3)
for base in (6, 13, 20, 27): rm(s, base, base + 5)  # card frames and inner rules
dark_chrome(s, 9)
for i, base in enumerate((6, 13, 20, 27)):
    circ, img, name, tag, desc = base + 1, base + 2, base + 3, base + 4, base + 6
    x = 0.65 + i * 3.1
    if i: vline(s, x - 0.2, 2.45, 4.0, WHITE, alpha=16)
    T(s, x, 2.45, 2.6, 0.3, '0%d' % (i + 1), 10.5, LAV, True)
    mv(sh(s, circ), x, 2.95, 0.85, 0.85)
    mv(sh(s, img), x + 0.24, 3.19, 0.37, 0.37)
    n = sh(s, name); mv(n, x, 4.05, 2.7, 0.78)
    n.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    for r in n.text_frame.paragraphs[0].runs: r.font.size = Pt(19)
    t = sh(s, tag); mv(t, x, 4.95, 2.6, 0.5)
    for r in t.text_frame.paragraphs[0].runs: r.font.size = Pt(11.5)
    d = sh(s, desc); mv(d, x, 5.55, 2.6, 1.0)
    for r in d.text_frame.paragraphs[0].runs: r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(LAV2)
notes(s, """
Nuestras cuatro soluciones de Inteligencia Artificial:

Agentic AI Studio — Estudio de contenidos con agentes de IA. Desde el ADN de tu marca, los agentes crean piezas multiformato listas para aprobar en minutos.
Agentic AI Reputation — Consultor de reputación. Asesora ante cualquier situación que afecte tu reputación con base histórica del negocio y la reputación.
GEO Rank — Plataforma de monitoreo GEO. Mide la visibilidad, las citas y la veracidad de tu marca en las IAs, y traza la ruta para posicionarte.
Data Pulse — Tracking de performance de embajadores digitales. Mide el desempeño de tus embajadores y convierte la data en acciones de comunicación.
""")

# ================================================================== 10 · FORMACIÓN (photo band + three tracks)
s = S[9]
keep_only(s, set())
pic(s, A + 'band10.jpg', 0, 0, 13.333, 3.1)
pic(s, A + 'fade_left.png', 0, 0, 13.333, 3.1)
pic(s, A + 'logo_white.png', 11.97, 0.42, 1.05, 0.37)
kicker(s, 0.65, 1.0, 8, 'Temario y casos a medida para cada empresa', LAV)
T(s, 0.65, 1.28, 8, 0.85, 'Formación', 44, WHITE, True)
T(s, 0.65, 2.2, 8, 0.45, 'Habilitamos a tu equipo para que trabaje con IA y lo haga por sí mismo.', 16, WHITE)
tracks = [('1', 'Reputación y liderazgo', 'Voceros creíbles', ['Marca Directiva y Embajadores']),
          ('2', 'Visibilidad y presencia digital', 'Prensa y entorno digital como aliados', ['Formación para periodistas', 'Political Intelligence AI']),
          ('3', 'Adopción y capacidades de IA', 'De la teoría al uso real de la IA', ['IA, data y automatización', 'IA para comunicación y marketing', 'Claude AI para empresas', 'Taller IA Agéntica'])]
for i, (n, frente, head, items) in enumerate(tracks):
    x = 0.65 + i * 4.18
    if i: vline(s, x - 0.25, 3.5, 2.75, 'E3E1EA')
    text(s, x, 3.45, 3.8, 0.3, [[(n + '  ', 12, PURPLE, True), (frente, 10.5, PURPLE, True)]])
    T(s, x, 3.8, 3.7, 0.75, head, 17, DARK, True, line=0.95)
    y = 4.65
    for it in items:
        hline(s, x, y, 3.7, 'E3E1EA')
        T(s, x, y + 0.08, 3.7, 0.3, it, 12, TXT, True, anchor='m')
        y += 0.4
text(s, 0.65, 6.5, 12, 0.3, [[('CLIENTES   ', 9.5, PURPLE, True, 300), ('BCP · Pacífico Seguros · Interbank · Antapaccay · Petrotal', 12, DARK, True)]])
notes(s, """
Temario y casos a medida para cada empresa. Habilitamos a tu equipo para que trabaje con IA y lo haga por sí mismo.

1. Reputación y liderazgo — Voceros creíbles
- Marca Directiva y Embajadores: líderes y colaboradores aprenden a comunicar con credibilidad en el entorno digital.

2. Visibilidad y presencia digital — Prensa y entorno digital como aliados
- Formación para periodistas: IA generativa aplicada al periodismo, con conferencia, taller práctico y certificado co-branding.
- Political Intelligence AI: exposiciones y talleres de una hora sobre coyuntura política.

3. Adopción y capacidades de IA — De la teoría al uso real de la IA
- IA, data y automatización: Copilot Studio y Power Platform aplicados a casos reales del área.
- IA para comunicación y marketing: IA generativa aplicada al análisis de data y a los flujos de contenido, campañas y ventas.
- Claude AI para empresas: implementación y onboarding del área para sistematizar el trabajo con IA en el día a día.
- Taller IA Agéntica: tu equipo aprende a crear y usar agentes de IA que ejecutan tareas de comunicación y marketing de principio a fin.

Clientes: BCP · Pacífico Seguros · Interbank · Antapaccay · Petrotal.
""")

# ================================================================== 11 · HERRAMIENTAS (light, open columns)
s = S[10]
rm(s, 2, 6, 17, 24, 31)
cols11 = [(0.95, [7, 8, 9, 10, 11, 12, 13, 14, 15, 16]), (4.01, [18, 19, 20, 21, 22, 23]),
          (7.08, [25, 26, 27, 28, 29, 30]), (10.14, [32, 33, 34, 35, 36])]
for i, (x0, ids) in enumerate(cols11):
    nx = 0.65 + i * 3.1
    for sid in ids:
        e = sh(s, sid); e.left = I(inch(e.left) - x0 + nx)
    if i: vline(s, nx - 0.2, 2.85, 3.85, 'E3E1EA')
mv(sh(s, 5), y=2.1)
notes(s, """
Trabajamos con los principales ecosistemas de IA y automatización:
- Ecosistema Microsoft: Copilot Studio, Power Apps, Power BI, Power Automate.
- Ecosistema Anthropic: Claude Projects, Claude Code.
- Ecosistema OpenAI: ChatGPT Projects, Codex.
- Ecosistema Google: Gemini Gems, Nano Banana.
""")

# ================================================================== 12 · RESULTADOS (dark, open data table)
s = S[11]
keep_only(s, {2, 4, 5, 6, 7, 20, 33})
dark_chrome(s, 12)
mv(sh(s, 6), y=2.08)
rows = [2.75, 4.17, 5.59]
for lab, y in zip((7, 20, 33), rows):
    mv(sh(s, lab), y=y, h=1.3)
for y in rows: hline(s, 0.65, y - 0.06, 12.03, WHITE, alpha=18)
cases = [
    (3.2, rows[0], 'PetroTal · Embajadores', [('+180%', 'interacciones'), ('+256%', 'publicaciones propias'), ('+13%', 'seguidores (+6,365)')]),
    (8.07, rows[0], 'Enel Perú · Marca Directiva', [('+2 pts', 'posicionamiento digital'), ('+68%', 'seguidores'), ('+35%', 'contenido propio'), ('+27%', 'interacciones')]),
    (3.2, rows[1], 'Promigas · Periodistas', [('146', 'periodistas'), ('+7,000', 'reproducciones'), ('1,500', 'interacciones'), (None, 'cobertura en medios')]),
    (8.07, rows[1], 'InRetail Pharma · Diagnóstico GEO', [('3', 'motores de IA auditados'), ('3', 'etapas de metodología'), ('2', 'frentes de escucha')]),
    (3.2, rows[2], 'BCP · Champions de IA', [('4.75/5', 'de satisfacción'), ('+140', 'colaboradores capacitados'), ('100%', 'aplicará lo aprendido')]),
]
for x0, y, title, metrics in cases:
    T(s, x0, y + 0.1, 4.62, 0.3, title, 11, LAV2, True)
    widths = {'Enel': [1.3, 1.1, 1.1, 1.12], 'Prom': [0.95, 1.3, 1.17, 1.2]}.get(title[:4], [1.54] * 3)
    big = 24 if len(metrics) == 4 else 26
    x = x0
    for (num, lbl), w in zip(metrics, widths):
        if num is None:
            T(s, x, y + 0.45, w - 0.05, 0.55, 'América TV, Infomercado', 10.5, WHITE, True, anchor='b')
        else:
            T(s, x, y + 0.42, w, 0.58, num, big, WHITE, True, anchor='b')
        T(s, x, y + 1.02, w - 0.1, 0.36, lbl, 9, LAV, False)
        x += w
vline(s, 7.85, rows[0] + 0.05, 2.75, WHITE, alpha=14)
notes(s, """
Resultados medidos con nuestros clientes: lo que logramos en cada frente.

Reputación y liderazgo
- PetroTal · Embajadores: +180% interacciones; +256% publicaciones propias; +13% seguidores (+6,365).
- Enel Perú · Marca Directiva (Marco Fragale): +2 pts posicionamiento digital; +68% seguidores; +35% contenido propio; +27% interacciones.

Visibilidad y presencia digital
- Promigas · Periodistas: 146 periodistas; +7,000 reproducciones; 1,500 interacciones; cobertura en medios (América TV, Infomercado).
- InRetail Pharma · Diagnóstico GEO: caso de diagnóstico. 3 motores de IA auditados (ChatGPT, Gemini, Perplexity); metodología en 3 etapas; 2 frentes de escucha.

Adopción y capacidades de IA
- BCP · Champions de IA: 4.75/5 de satisfacción; +140 colaboradores capacitados; 100% aplicará lo aprendido.
""")

# ================================================================== 13 · EQUIPO (light, portraits)
s = S[12]
rm(s, 25, 10, 20)
mv(sh(s, 5), w=12.03)
bios = {
    6: ['+20 años de experiencia en Investigación de Mercados, Marketing Digital, Gestión de Marcas Directivas y Reputación.',
        'Maestría en Marketing y Transformación Digital por INCAE Business School.',
        'Docente de Postgrado en Universidad de Piura y programas corporativos de UTEC. Columnista en Diario Gestión.'],
    16: ['Especialista en Marketing e Inteligencia Artificial Generativa.',
         'Maestría en Marketing y Ventas por INCAE Business School.',
         'Docente de IA Generativa en UDEP, CENTRUM PUCP, UPC y USIL.'],
}
for photo, name, role, li in [(6, 7, 8, 9), (16, 17, 18, 19)]:
    ph = sh(s, photo); mv(ph, y=2.7, w=1.6, h=1.6)
    bx = inch(ph.left); px = bx + 1.85
    n = sh(s, name); mv(n, x=px, y=2.88, h=0.5)
    for r in n.text_frame.paragraphs[0].runs: r.font.size = Pt(24)
    mv(sh(s, role), x=px, y=3.42); mv(sh(s, li), x=px, y=3.75)
    text(s, bx, 4.55, 5.77, 1.65, [[(b, 12, TXT, False)] for b in bios[photo]], after=5, bullet=PURPLE)
notes(s, """
Nosotros te acompañamos. Potenciamos la comunicación digital en empresas, mejorando el posicionamiento y reputación de directivos y equipos.

Víctor Lozano — Director de Innovación (linkedin.com/in/victorlozanou)
- +20 años de experiencia en Investigación de Mercados, Marketing Digital, Gestión de Marcas Directivas y Reputación.
- Ha liderado estrategias de posicionamiento corporativo en Telecom, Banca, Consumo Masivo, Minería y Retail.
- Maestría en Marketing y Transformación Digital por INCAE Business School.
- Docente de Postgrado en Universidad de Piura y programas corporativos de UTEC. Columnista en Diario Gestión.

Julio Talledo — Director Comercial & Marketing (linkedin.com/in/juliotalledo)
- Especialista en Marketing e Inteligencia Artificial Generativa. Director Comercial y de Marketing de Verne.
- Maestría en Marketing y Ventas por INCAE Business School.
- Docente de IA Generativa en UDEP, CENTRUM PUCP, UPC y USIL.
- Ha colaborado con BCP, Grupo Romero, Universidad de Piura, Ransa, LAP, Bitel y más.

www.verne.la
""")

for sl in S:
    el = sl._element
    tr = el.makeelement(qn('p:transition'), {'spd': 'med'})
    tr.append(tr.makeelement(qn('p:fade'), {}))
    el.find(qn('p:clrMapOvr')).addnext(tr) if el.find(qn('p:clrMapOvr')) is not None else el.find(qn('p:cSld')).addnext(tr)
p.save(OUT)

# The source deck declares jpg as the non-standard "image/jpg"; PowerPoint expects "image/jpeg".
tmp = OUT + '.tmp'
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
    for item in zi.infolist():
        data = zi.read(item.filename)
        if item.filename == '[Content_Types].xml':
            data = data.replace(b'ContentType="image/jpg"', b'ContentType="image/jpeg"')
            if b'Extension="jpg"' not in data:
                data = data.replace(b'<Default Extension="png"', b'<Default Extension="jpg" ContentType="image/jpeg"/><Default Extension="png"')
        zo.writestr(item, data)
shutil.move(tmp, OUT)
print('saved', OUT)

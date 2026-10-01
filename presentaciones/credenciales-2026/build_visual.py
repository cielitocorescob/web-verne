import copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC, OUT = 'Credenciales_VERNE_2026.pptx', 'Credenciales_VERNE_2026_Visual.pptx'
p = Presentation(SRC)
S = p.slides

PURPLE, DARK, LAV, LAV2, LAVBG, GRAY, TXT, WHITE = '6D4ABE', '25283D', 'C0B0F4', 'D8D4EA', 'F3F0FB', '6B6E80', '333333', 'FFFFFF'


def I(v):
    return Emu(int(round(v * 914400)))


def sh(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)


def rm(slide, *ids):
    for sid in ids:
        e = sh(slide, sid)._element
        e.getparent().remove(e)


def mv(shape, x=None, y=None, w=None, h=None):
    if x is not None: shape.left = I(x)
    if y is not None: shape.top = I(y)
    if w is not None: shape.width = I(w)
    if h is not None: shape.height = I(h)


def size(shape, pt):
    for par in shape.text_frame.paragraphs:
        for r in par.runs:
            r.font.size = Pt(pt)


def settext(shape, text, run=0):
    """Replace the text of one run, dropping the paragraph's other runs (keeps formatting)."""
    par = shape.text_frame.paragraphs[0]
    runs = par.runs
    runs[run].text = text
    for i, r in enumerate(runs):
        if i != run:
            r._r.getparent().remove(r._r)
    for extra in shape.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def _font(r, pt, color, bold):
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.makeelement(qn(tag), {'typeface': 'Poppins'})
        rPr.append(el)


def text(slide, x, y, w, h, paras, align='l', anchor='t', spacing=None):
    """paras: list of paragraphs; each paragraph a list of (text, pt, color, bold) runs."""
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    for i, runs in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if spacing:
            par.space_after = Pt(spacing)
        for t, pt, color, bold in runs:
            r = par.add_run()
            r.text = t
            _font(r, pt, color, bold)
    return tb


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None:
        s.adjustments[0] = radius
    s.text_frame.text = ''
    return s


def chip(slide, x, y, w, h, label, fill, color, pt=11, bold=True):
    box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(slide, x, y, w, h, [[(label, pt, color, bold)]], align='c', anchor='m')


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


# ---------------------------------------------------------------- global clean-up
# Slide-level page numbers duplicate the master's ‹Nº› (bottom-left) and collide with the
# master's "www.verne.la" (bottom-right). Keep them only where a full-bleed image hides the master.
for n, sid in [(3, 49), (4, 24), (5, 34), (7, 18), (8, 21), (10, 37), (11, 37)]:
    rm(S[n - 1], sid)
# Top-left wordmark duplicates the master's logo on light slides.
for n in (4, 5, 7, 8, 10, 11):
    rm(S[n - 1], 2 if n != 5 else 3)

# ---------------------------------------------------------------- 1 · Portada
s = S[0]
t = sh(s, 6)
mv(t, 0.65, 2.35, 6.6, 2.2)
size(t, 32)
mark = copy.deepcopy(sh(S[8], 3)._element)  # white "verne · Comunicación e Innovación" wordmark from the dark slides
s.shapes._spTree.append(mark)
notes(s, """
Apertura. VERNE: Consultora Estratégica de Comunicación Digital, potenciada con Inteligencia Artificial.
www.verne.la
""")

# ---------------------------------------------------------------- 2 · Quiénes somos
s = S[1]
rm(s, 5, 12, 21, 30, 13, 22, 31, 35)
for col, (c, cn, lab, stat, desc, src) in enumerate([(6, 7, 8, 9, 10, 11), (15, 16, 17, 18, 19, 20), (24, 25, 26, 27, 28, 29)]):
    for i in (c, cn, lab):
        mv(sh(s, i), y=2.55)
    st = sh(s, stat); mv(st, y=3.05, h=1.35); size(st, 76)
    d = sh(s, desc); mv(d, y=4.5, h=0.8); size(d, 12.5)
    mv(sh(s, src), y=5.45)
for v in (14, 23):
    mv(sh(s, v), y=2.55, h=3.2)
notes(s, """
La comunicación estratégica ya no vive solo en los medios. Hoy se juega en tres frentes:

1. Reputación y liderazgo — 3x más confianza cuando la información la comparte un directivo o colaborador que la cuenta institucional (Fuente: Edelman Trust Barometer).
Posicionamos a la marca corporativa, a sus líderes y embajadores en el ecosistema digital.

2. Visibilidad y presencia digital — 70% de los consumidores peruanos usa IA para informarse antes de comprar o contratar (Fuente: Impronta Research, ago 2025).
Logramos que te encuentren, citen y recomienden en buscadores, IA y medios.

3. Adopción y capacidades de IA — 90% de las áreas de Comunicación y Marketing en el mundo ya experimenta con IA (Fuente: McKinsey, 2026).
Habilitamos a los equipos para usar IA generativa en su día a día.

Cierre: Ni solo tecnología ni solo comunicación tradicional.
""")

# ---------------------------------------------------------------- 3 · Clientes (sin cambios de contenido)
notes(S[2], """
Confían en nosotros: empresas líderes de sectores regulados y competitivos.
Sectores: Energía, Minería e Hidrocarburos · Banca y Servicios Financieros · Consumo, Retail e Industria · Educación, Servicios e Infraestructura.
""")

# ---------------------------------------------------------------- 4 · Servicios (sin cambios de contenido)
notes(S[3], """
Nuestros servicios se organizan en dos modalidades para cada uno de los tres frentes:
- CONSULTORÍA (lo hacemos por ti)
- FORMACIÓN (habilitamos a tu equipo)

1. Reputación y liderazgo — Consultoría: Marca Directiva y Corporativa; Embajadores digitales. Formación: Talleres de Marca Directiva y Embajadores.
2. Visibilidad y presencia digital — Consultoría: Posicionamiento GEO; Marketing de Contenidos; Asuntos Públicos. Formación: Formación para periodistas; Political Intelligence AI.
3. Adopción y capacidades de IA — Consultoría: Diagnóstico de madurez digital; Agentic AI Studio; Agentic AI Reputation; GEO Rank; Data Pulse. Formación: Talleres de IA, data y automatización; Taller de IA para comunicación y MKT; Claude AI para empresas; Taller IA Agéntica.
""")

# ---------------------------------------------------------------- 5 · Marca Corporativa y Directiva
s = S[4]
lead = sh(s, 6)
settext(lead, 'Los líderes cambian, la marca permanece.')
mv(lead, y=2.12, h=0.5)
lead.text_frame.paragraphs[0].runs[0].font.bold = True
lead.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(PURPLE)
size(lead, 20)
for rect, ttl, desc, phrase in [(7, 8, 9, 'Es lo que permanece.'), (10, 11, 12, 'Es lo que la amplifica.')]:
    mv(sh(s, rect), y=2.85, h=1.2)
    mv(sh(s, ttl), y=3.03, h=0.4); size(sh(s, ttl), 18)
    d = sh(s, desc); settext(d, phrase); mv(d, y=3.5, h=0.4); size(d, 14)
rm(s, 15, 18, 21)
for line, lab in [(13, 14), (16, 17), (19, 20)]:
    mv(sh(s, line), y=4.45)
    l = sh(s, lab); mv(l, y=4.62, h=0.5); size(l, 18)
notes(s, """
Construimos la marca de tu empresa a través de la voz de sus líderes. Los líderes cambian, la marca permanece.

Marca Corporativa: la reputación y credibilidad de la empresa en el ecosistema digital. Es lo que permanece.
Marca Directiva: tus C-level, directivos y voceros dan rostro y voz a esa narrativa. Es lo que la amplifica.

Cómo trabajamos:
01 Diagnóstico — Analizamos la presencia digital de la empresa y de sus directivos, y la comparamos con la de sus pares en la industria.
02 Estrategia — Definimos la narrativa corporativa, los ejes de contenido y el rol de cada vocero, para que el mensaje sea de la empresa y no dependa de una sola persona.
03 Gestión — Creamos contenidos, acompañamos a los voceros y medimos el desempeño de la marca corporativa y de sus líderes.
""")

# ---------------------------------------------------------------- 6 · Embajadores Digitales
s = S[5]
sub = sh(s, 6); size(sub, 16); mv(sub, y=2.1, w=8.08, h=0.85)
rm(s, 9, 13, 17)
for num, lab in [(7, 8), (11, 12), (15, 16)]:
    n = sh(s, num); mv(n, y=3.45, h=1.0); size(n, 60)
    l = sh(s, lab); mv(l, y=4.5, h=0.45); size(l, 18)
for v in (10, 14):
    mv(sh(s, v), y=3.55, h=1.45)
mv(sh(s, 26), x=4.60, y=7.0)
notes(s, """
Amplificamos la comunicación corporativa a través de tus colaboradores, quienes hablan como la voz más genuina de la marca.

01 Estrategia — Seleccionamos y activamos embajadores alineados a los objetivos de la marca.
02 Acompañamiento — Buenas prácticas, benchmark y co-creación de contenidos, con soporte continuo.
03 Medición — Data pública, matriz de evaluación e informe mensual de desempeño, con Data Pulse.
""")

# ---------------------------------------------------------------- 7 · Posicionamiento GEO
s = S[6]
lead = sh(s, 5)
settext(lead, 'Logramos que la IA cite y recomiende a tu marca.')
lead.text_frame.paragraphs[0].runs[0].font.bold = True
lead.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(DARK)
size(lead, 20); mv(lead, y=2.1, h=0.5)
x = 0.65
for name, w in [('ChatGPT', 1.2), ('Google AI Overview', 2.1), ('Perplexity', 1.35), ('Claude', 1.0)]:
    chip(s, x, 2.78, w, 0.4, name, LAVBG, PURPLE, 11)
    x += w + 0.15
rm(s, 8, 11, 14)
# Gantt-style timeline: bar widths proportional to 15 / 15 / 60 days.
gap, total = 0.06, 12.03
unit = (total - 2 * gap) / 90
x = 0.65
for bar, lab, days, name in [(6, 7, 15, 'Diagnóstico'), (9, 10, 15, 'Estrategia'), (12, 13, 60, 'Implementación')]:
    w = days * unit
    mv(sh(s, bar), x, 3.62, w, 0.34)
    l = sh(s, lab); settext(l, name, run=1); mv(l, x, 4.12, w, 0.36)
    l.text_frame.paragraphs[0].runs[0].text = name  # rebuilt below with number
    l.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    x += w + gap
# restore "0N  " prefix as a separate purple run
for lab, num in [(7, '01  '), (10, '02  '), (13, '03  ')]:
    l = sh(s, lab)
    par = l.text_frame.paragraphs[0]
    r0 = par.runs[0]
    new = copy.deepcopy(r0._r)
    r0._r.addprevious(new)
    par.runs[0].text = num
    par.runs[0].font.color.rgb = RGBColor.from_string(PURPLE)
    par.runs[1].font.color.rgb = RGBColor.from_string(DARK)
    size(l, 13)
x = 0.65
for days in (15, 15, 60):
    w = days * unit
    text(s, x, 4.48, w, 0.6, [[(str(days), 26, PURPLE, True), (' días', 13, GRAY, False)]])
    x += w + gap
notes(s, """
Logramos que la IA cite y recomiende a tu marca. Trabajamos para que ChatGPT, Google AI Overview, Perplexity y Claude te mencionen en tu categoría.

01 Diagnóstico (15 días) — Una radiografía de cómo aparece tu marca frente a la competencia, con brechas y oportunidades.
02 Estrategia (15 días) — Definimos narrativas, arquitectura de contenidos y el recorrido del usuario en la IA.
03 Implementación (60 días) — Parrilla GEO, producción de contenidos y verificación de citación en las IAs.

Resultados medibles en visibilidad, posición, citaciones y share of voice en IA.
Clientes: Gloria · ISA.
""")

# ---------------------------------------------------------------- 8 · Diagnóstico de Madurez Digital
s = S[7]
rm(s, 12, 16, 20)
for rect, num, ttl in [(9, 10, 11), (13, 14, 15), (17, 18, 19)]:
    r = sh(s, rect)
    y0 = r.top / 914400
    t = sh(s, ttl); mv(t, y=y0 + 0.3, h=0.42); size(t, 17)
cy = sh(s, 9).top / 914400 + 0.88
text(s, 7.25, cy, 1.9, 0.42, [[('Encuesta digital', 12, PURPLE, True)]], anchor='m')
text(s, 8.8, cy, 0.3, 0.42, [[('+', 16, GRAY, True)]], align='c', anchor='m')
text(s, 9.15, cy, 2.9, 0.42, [[('Entrevistas a profundidad', 12, PURPLE, True)]], anchor='m')
cy = sh(s, 13).top / 914400 + 0.78
text(s, 7.25, cy, 1.9, 0.6, [[('5', 30, PURPLE, True), (' niveles', 12, GRAY, False)]], anchor='b')
text(s, 8.95, cy, 2.2, 0.6, [[('7', 30, PURPLE, True), (' indicadores', 12, GRAY, False)]], anchor='b')
for i, c in enumerate(['E3E1EA', 'D8D4EA', LAV, '8F72D6', PURPLE]):  # 5-level maturity scale
    h = 0.14 + i * 0.11
    box(s, 11.25 + i * 0.24, cy + 0.6 - h, 0.18, h, c)
cy = sh(s, 17).top / 914400 + 0.86
text(s, 7.25, cy, 1.35, 0.4, [[('Hoja de ruta', 12, GRAY, False)]], anchor='m')
x = 8.6
for lbl, fill, col in [('0-3', LAV, DARK), ('3-9', PURPLE, WHITE), ('9-18', DARK, WHITE)]:
    chip(s, x, cy, 0.95, 0.4, lbl, fill, col, 12)
    x += 1.03
text(s, x + 0.05, cy, 0.8, 0.4, [[('meses', 12, GRAY, False)]], anchor='m')
notes(s, """
Te decimos qué tan bien usan tus equipos la IA y dónde invertir primero. Los niveles se inspiran en el modelo CMMI.

01 Medimos la adopción — A través de una encuesta digital representativa y entrevistas a profundidad con quienes deciden, habilitan y construyen soluciones.
02 Comparamos — Ubicamos a cada área en una escala de 5 niveles según su uso de IA generativa, automatización y datos. Cada una se mide con 7 indicadores, como autonomía, impacto y capacitación.
03 Priorizamos — Dashboard interactivo, informe ejecutivo, segmentación de gerencias y hoja de ruta a 0-3, 3-9 y 9-18 meses.

Clientes: BCP / Credicorp · Primax.
""")

# ---------------------------------------------------------------- 9 · Cuatro soluciones de IA
s = S[8]
for base in (6, 13, 20, 27):
    rect, circ, img, name, tag, line, desc = range(base, base + 7)
    rm(s, desc)
    mv(sh(s, rect), y=2.5, h=3.9)
    c = sh(s, circ); mv(c, y=2.95, w=0.8, h=0.8)
    cx = c.left / 914400
    mv(sh(s, img), cx + 0.22, 2.95 + 0.22, 0.36, 0.36)
    n = sh(s, name); mv(n, y=4.2, h=0.7); size(n, 18)
    mv(sh(s, line), y=5.05)
    t = sh(s, tag); mv(t, y=5.22, h=0.7); size(t, 12)
notes(s, """
Nuestras cuatro soluciones de Inteligencia Artificial:

Agentic AI Studio — Estudio de contenidos con agentes de IA. Desde el ADN de tu marca, los agentes crean piezas multiformato listas para aprobar en minutos.
Agentic AI Reputation — Consultor de reputación. Asesora ante cualquier situación que afecte tu reputación con base histórica del negocio y la reputación.
GEO Rank — Plataforma de monitoreo GEO. Mide la visibilidad, las citas y la veracidad de tu marca en las IAs, y traza la ruta para posicionarte.
Data Pulse — Tracking de performance de embajadores digitales. Mide el desempeño de tus embajadores y convierte la data en acciones de comunicación.
""")

# ---------------------------------------------------------------- 10 · Formación
s = S[9]
rm(s, 11, 17, 20, 26, 29, 32, 35)
for rect in (6, 12, 21):
    mv(sh(s, rect), y=2.62, h=3.2)
for hl in (8, 14, 23):
    h = sh(s, hl); mv(h, y=3.22, h=0.9); size(h, 19)
for line, item, y in [(9, 10, 4.3), (15, 16, 4.3), (18, 19, 5.0), (24, 25, 4.3), (27, 28, 4.3), (30, 31, 5.0), (33, 34, 5.0)]:
    mv(sh(s, line), y=y)
    it = sh(s, item); mv(it, y=y + 0.12, h=0.5); size(it, 13)
cl = sh(s, 36)
mv(cl, y=6.15)
for r in cl.text_frame.paragraphs[0].runs:
    r.text = r.text.replace(' (logos)', '')
notes(s, """
Temario y casos a medida para cada empresa. Habilitamos a tu equipo para que trabaje con IA y lo haga por sí mismo.

1. Reputación y liderazgo — Voceros creíbles
- Marca Directiva y Embajadores: líderes y colaboradores aprenden a comunicar con credibilidad en el entorno digital.

2. Visibilidad y presencia digital — Prensa y entorno digital como aliados
- Formación para periodistas: IA generativa aplicada al periodismo, con conferencia, taller práctico y certificado co-branding.
- Political Intelligence AI: exposiciones y talleres de una hora sobre coyuntura política.

3. Adopción y capacidades de IA — De la teoría al uso real de la IA
- IA, data y automatización: Copilot Studio y Power Platform aplicados a casos reales del área.
- IA para comunicación y marketing: IA generativa aplicada al análisis de data y a los flujos de contenido, campañas y ventas.
- Claude AI para empresas: implementación y onboarding del área para sistematizar el trabajo con IA en el día a día.
- Taller IA Agéntica: tu equipo aprende a crear y usar agentes de IA que ejecutan tareas de comunicación y marketing de principio a fin.

Clientes: BCP · Pacífico Seguros · Interbank · Antapaccay · Petrotal.
""")

# ---------------------------------------------------------------- 11 · Herramientas (sin cambios de contenido)
notes(S[10], """
Trabajamos con los principales ecosistemas de IA y automatización:
- Ecosistema Microsoft: Copilot Studio, Power Apps, Power BI, Power Automate.
- Ecosistema Anthropic: Claude Projects, Claude Code.
- Ecosistema OpenAI: ChatGPT Projects, Codex.
- Ecosistema Google: Gemini Gems, Nano Banana.
""")

# ---------------------------------------------------------------- 12 · Casos reales
s = S[11]
rm(s, 6, *range(10, 13), *range(15, 19), *range(23, 27), *range(29, 32), *range(36, 39))
settext(sh(s, 14), 'Enel Perú · Marca Directiva')
rows = [2.35, 3.88, 5.41]
for lab, y in zip((7, 20, 33), rows):
    mv(sh(s, lab), y=y, h=1.35)
for line, y in zip((19, 32), rows[1:]):
    mv(sh(s, line), y=y - 0.08)
cards = [
    (8, 9, rows[0], [('+180%', 'interacciones'), ('+256%', 'publicaciones propias'), ('+13%', 'seguidores (+6,365)')]),
    (13, 14, rows[0], [('+2 pts', 'posicionamiento digital'), ('+68%', 'seguidores'), ('+35%', 'contenido propio'), ('+27%', 'interacciones')]),
    (21, 22, rows[1], [('146', 'periodistas'), ('+7,000', 'reproducciones'), ('1,500', 'interacciones')]),
    (27, 28, rows[1], [('3', 'motores de IA auditados'), ('3', 'etapas de metodología'), ('2', 'frentes de escucha')]),
    (34, 35, rows[2], [('4.75/5', 'de satisfacción'), ('+140', 'colaboradores capacitados'), ('100%', 'aplicará lo aprendido')]),
]
for rect, ttl, y, metrics in cards:
    r = sh(s, rect); mv(r, y=y, h=1.35)
    t = sh(s, ttl); mv(t, y=y + 0.12, h=0.3)
    x0 = r.left / 914400 + 0.22
    inner = r.width / 914400 - 0.44
    widths = [1.25, 0.99, 0.99, 0.99] if len(metrics) == 4 else [inner / 3] * 3
    x = x0
    for (num, lbl), w in zip(metrics, widths):
        text(s, x, y + 0.45, w, 0.45, [[(num, 22, WHITE, True)]], anchor='b')
        text(s, x, y + 0.92, w - 0.08, 0.38, [[(lbl, 9, LAV, False)]])
        x += w
notes(s, """
Resultados medidos con nuestros clientes: lo que logramos en cada frente.

Reputación y liderazgo
- PetroTal · Embajadores: +180% interacciones; +256% publicaciones propias; +13% seguidores (+6,365).
- Enel Perú · Marca Directiva (Marco Fragale): +2 pts posicionamiento digital; +68% seguidores; +35% contenido propio; +27% interacciones.

Visibilidad y presencia digital
- Promigas · Periodistas: 146 periodistas; +7,000 reproducciones; 1,500 interacciones; cobertura en medios (América TV, Infomercado).
- InRetail Pharma · Diagnóstico GEO: caso de diagnóstico. 3 motores de IA auditados (ChatGPT, Gemini, Perplexity); metodología en 3 etapas; 2 frentes de escucha.

Adopción y capacidades de IA
- BCP · Champions de IA: 4.75/5 de satisfacción; +140 colaboradores capacitados; 100% aplicará lo aprendido.
""")

# ---------------------------------------------------------------- 13 · Equipo
s = S[12]
rm(s, 25, 26)
mv(sh(s, 5), w=12.03)
for photo, name, role, li, bio, hi in [
    (6, 7, 8, 9, 10, [('+20 años', PURPLE, True), (' de experiencia en Investigación de Mercados, Marketing Digital, Gestión de Marcas Directivas y Reputación.', DARK, False)]),
    (16, 17, 18, 19, 20, [('Especialista en ', DARK, False), ('Marketing e Inteligencia Artificial Generativa.', PURPLE, True)]),
]:
    ph = sh(s, photo); mv(ph, y=2.75, w=1.5, h=1.5)
    px = ph.left / 914400 + 1.75
    mv(sh(s, name), x=px, y=2.95); size(sh(s, name), 22)
    mv(sh(s, role), x=px, y=3.42)
    mv(sh(s, li), x=px, y=3.74)
    b = sh(s, bio)
    bx, bw = b.left / 914400, b.width / 914400
    rm(s, bio)
    text(s, bx, 4.6, bw, 1.3, [[(t_, 16, c_, bold_) for t_, c_, bold_ in hi]])
notes(s, """
Nosotros te acompañamos. Potenciamos la comunicación digital en empresas, mejorando el posicionamiento y reputación de directivos y equipos.

Víctor Lozano — Director de Innovación (linkedin.com/in/victorlozanou)
- +20 años de experiencia en Investigación de Mercados, Marketing Digital, Gestión de Marcas Directivas y Reputación.
- Ha liderado estrategias de posicionamiento corporativo en Telecom, Banca, Consumo Masivo, Minería y Retail.
- Maestría en Marketing y Transformación Digital por INCAE Business School.
- Docente de Postgrado en Universidad de Piura y programas corporativos de UTEC. Columnista en Diario Gestión.

Julio Talledo — Director Comercial & Marketing (linkedin.com/in/juliotalledo)
- Especialista en Marketing e Inteligencia Artificial Generativa. Director Comercial y de Marketing de Verne.
- Maestría en Marketing y Ventas por INCAE Business School.
- Docente de IA Generativa en UDEP, CENTRUM PUCP, UPC y USIL.
- Ha colaborado con BCP, Grupo Romero, Universidad de Piura, Ransa, LAP, Bitel y más.

www.verne.la
""")

p.save(OUT)
print('saved', OUT)

# The source deck declares jpg as the non-standard "image/jpg"; PowerPoint expects "image/jpeg".
import zipfile, shutil
tmp = OUT + '.tmp'
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
    for item in zi.infolist():
        data = zi.read(item.filename)
        if item.filename == '[Content_Types].xml':
            data = data.replace(b'ContentType="image/jpg"', b'ContentType="image/jpeg"')
            if b'Extension="jpg"' not in data:
                data = data.replace(b'<Default Extension="png"', b'<Default Extension="jpg" ContentType="image/jpeg"/><Default Extension="png"')
        zo.writestr(item, data)
shutil.move(tmp, OUT)

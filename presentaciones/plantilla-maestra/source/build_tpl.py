"""Builds the VERNE master template (.potx) and an identical .pptx."""
import copy, io, re, zipfile, shutil
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT, CONTENT_TYPE as CT
from pptx.opc.packuri import PackURI
from pptx.parts.slide import SlideLayoutPart
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

import tpl_layouts as T
from tpl_layouts import (PURPLE, NAVY, LAV, TXT, WHITE, LAVBG, LINE, GRAY, NOTE, MID, PANEL, DEEP, M, R, CW)

MEDIA = '/tmp/claude-0/-home-user-web-verne/7d4b3b84-7c70-578c-b33c-1b2bd060621f/scratchpad/u/ppt/media/'
A, IC = 'assets/', 'icons/out/'
OUT = 'VERNE_Plantilla_Maestra'
IMG = {'logo': MEDIA + 'image1.png', 'logo_white': A + 'logo_white.png',
       'glow_a': MEDIA + 'image36.jpg', 'glow_b': MEDIA + 'image50.jpg'}
PHOTO_AUDIENCE, PHOTO_MAN, PHOTO_WOMAN = MEDIA + 'image2.jpg', MEDIA + 'image34.jpg', MEDIA + 'image35.jpg'

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
master_part = prs.slide_master.part

# ------------------------------------------------------------------ theme: colors + Poppins
theme_part = master_part.part_related_by(RT.THEME)
th = etree.fromstring(theme_part.blob)
ans = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
th.set('name', 'VERNE')
cs = th.find('.//a:clrScheme', ans); cs.set('name', 'VERNE')
for tag, val in T.THEME_COLORS:
    el = cs.find(f'a:{tag}', ans)
    for c in list(el): el.remove(c)
    etree.SubElement(el, qn('a:srgbClr')).set('val', val)
fs = th.find('.//a:fontScheme', ans); fs.set('name', 'VERNE · Poppins')
for grp in ('a:majorFont', 'a:minorFont'):
    fs.find(grp, ans).find('a:latin', ans).set('typeface', 'Poppins')
esl = th.find('.//a:effectStyleLst', ans)
for es in list(esl): esl.remove(es)
for _ in range(3):
    etree.SubElement(etree.SubElement(esl, qn('a:effectStyle')), qn('a:effectLst'))
theme_part._blob = etree.tostring(th, xml_declaration=True, encoding='UTF-8', standalone=True)

# ------------------------------------------------------------------ layouts + master
old = [rId for rId, rel in master_part.rels.items() if rel.reltype == RT.SLIDE_LAYOUT]
for rId in old:
    master_part.drop_rel(rId)
layouts = T.build_layouts()
new_rids = []
for n, L in enumerate(layouts, 1):
    lp = SlideLayoutPart(PackURI(f'/ppt/slideLayouts/slideLayout{n}.xml'), CT.PML_SLIDE_LAYOUT, prs.part.package,
                         parse_xml('<p:sldLayout xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'))
    xml = L.xml()
    for key in L.images:
        _, rid = lp.get_or_add_image_part(IMG[key])
        xml = xml.replace('{' + key + '}', rid)
    lp._element = parse_xml(xml.encode('utf-8'))
    lp.relate_to(master_part, RT.SLIDE_MASTER)
    new_rids.append(master_part.relate_to(lp, RT.SLIDE_LAYOUT))
_, logo_rid = master_part.get_or_add_image_part(IMG['logo'])
mx = T.master_xml(new_rids, logo_rid).replace('<p:ph idx="1"', '<p:ph type="body" idx="1"')
master_part._element = parse_xml(mx.encode('utf-8'))
master_part.__dict__.pop('slide_master', None)  # drop python-pptx's cached wrapper of the old master

# ------------------------------------------------------------------ table style + editing guides
prs.part.part_related_by(RT.TABLE_STYLES)._blob = T.table_styles_xml().encode('utf-8')
g = lambda inch: int(round(inch * 576))
guides = ''.join(f'<p:guide orient="vert" pos="{g(v)}"/>' for v in (M, R, 13.333 / 2)) + \
         ''.join(f'<p:guide pos="{g(v)}"/>' for v in (1.0, 6.85))
prs.part.part_related_by(RT.VIEW_PROPS)._blob = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:viewPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:normalViewPr><p:restoredLeft sz="15620"/><p:restoredTop sz="94660"/></p:normalViewPr><p:slideViewPr><p:cSldViewPr snapToGrid="0">'
    '<p:cViewPr varScale="1"><p:scale><a:sx n="100" d="100"/><a:sy n="100" d="100"/></p:scale><p:origin x="0" y="0"/></p:cViewPr>'
    f'<p:guideLst>{guides}</p:guideLst></p:cSldViewPr></p:slideViewPr><p:gridSpacing cx="76200" cy="76200"/></p:viewPr>').encode('utf-8')

LY = {l.name[:2]: l for l in prs.slide_layouts}


# ------------------------------------------------------------------ slide helpers
def I(v): return Emu(int(round(v * 914400)))


def new(code, notes=None):
    s = prs.slides.add_slide(LY[code])
    if notes: s.notes_slide.notes_text_frame.text = notes.strip()
    return s


def fill(s, idx, value):
    """Fill a placeholder (idx int, or 'title'). value: str, or list of paragraphs [(text, level)]."""
    ph = s.shapes.title if idx == 'title' else s.placeholders[idx]
    tf = ph.text_frame
    if isinstance(value, str):
        value = [(value, 0)]
    for i, (t, lvl) in enumerate(value):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = t; p.level = lvl
    return ph


def drop(s, *idxs):
    for i in idxs:
        e = s.placeholders[i]._element
        e.getparent().remove(e)


def photo(s, idx, path):
    return s.placeholders[idx].insert_picture(path)


def run_fmt(r, sz, color, b=False, i=False, spc=None):
    r.font.size = Pt(sz); r.font.bold = b; r.font.italic = i
    r.font.color.rgb = RGBColor.from_string(color); r.font.name = 'Poppins'
    if spc: r._r.get_or_add_rPr().set('spc', str(spc))


def tx(s, x, y, w, h, paras, align='l', anchor='t', after=0, line=None):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, tuple): paras = [[paras]]
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if after: p.space_after = Pt(after)
        if line: p.line_spacing = line
        for t, *f in runs:
            r = p.add_run(); r.text = t; run_fmt(r, *f)
    return tb


def shape(s, x, y, w, h, color, kind=MSO_SHAPE.RECTANGLE, line=None, alpha=None, radius=None):
    sp = s.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    if color:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(color)
        if alpha is not None:
            clr = sp._element.spPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            clr.append(clr.makeelement(qn('a:alpha'), {'val': str(alpha * 1000)}))
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(0.75)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if radius is not None: sp.adjustments[0] = radius
    return sp


def hl(s, x, y, w, color=LINE): return shape(s, x, y, w, 0.012, color)
def vl(s, x, y, h, color=LINE): return shape(s, x, y, 0.012, h, color)


def chip(s, x, y, w, label, fill, color, h=0.34, line=None):
    shape(s, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, line=line, radius=0.5)
    tx(s, x, y, w, h, [[(label, 10, color, True)]], align='c', anchor='m')


def cover_pic(s, path, x, y, w, h, crop_bias=0.5):
    """Picture cropped to fill the box (object-fit: cover), stays an editable picture."""
    im = Image.open(path); iw, ih = im.size
    pic = s.shapes.add_picture(path, I(x), I(y), I(w), I(h))
    box_r, img_r = w / h, iw / ih
    if img_r > box_r:
        c = (1 - box_r / img_r); pic.crop_left = c * crop_bias; pic.crop_right = c * (1 - crop_bias)
    else:
        c = (1 - img_r / box_r); pic.crop_top = c * crop_bias; pic.crop_bottom = c * (1 - crop_bias)
    return pic


def eyebrow(s, x, y, w, t, color=PURPLE): return tx(s, x, y, w, 0.28, [[(t.upper(), 10.5, color, True, False, 300)]])


def icon_disc(name, size=256, ring=PURPLE):
    """Purple disc with a white line icon (for icon placeholders)."""
    disc = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    from PIL import ImageDraw
    ImageDraw.Draw(disc).ellipse((0, 0, size - 1, size - 1), fill=tuple(int(ring[i:i + 2], 16) for i in (0, 2, 4)) + (255,))
    ic = Image.open(f'{IC}{name}_w.png').resize((int(size * 0.46),) * 2)
    o = (size - ic.size[0]) // 2
    disc.alpha_composite(ic, (o, o))
    path = f'{A}disc_{name}.png'; disc.save(path); return path


# ======================================================================================
# GUÍA RÁPIDA
# ======================================================================================
s = new('01', 'Portada de la guía. Diseño 01 Portada principal.')
fill(s, 10, 'SISTEMA DE PRESENTACIONES')
fill(s, 'title', 'Plantilla maestra VERNE')
fill(s, 1, 'Guía de uso, recursos y catálogo de diseños para crear presentaciones con la identidad VERNE.')
fill(s, 11, 'Versión 1.0 · 2026')

s = new('03', 'Diseño 03 Apertura de sección.')
fill(s, 12, '01'); fill(s, 'title', 'Guía rápida'); fill(s, 11, 'Lo esencial para usar la plantilla sin romper el sistema visual.')

s = new('13', 'Diseño 13 Proceso / pasos.')
fill(s, 10, 'GUÍA RÁPIDA · CÓMO USAR'); fill(s, 'title', 'Crea una presentación en cuatro pasos')
fill(s, 11, 'Todo el estilo vive en el Patrón de diapositivas: tú solo eliges el diseño y escribes.')
for i, (h, b) in enumerate([
        ('Elige un diseño', 'Inicio › Nueva diapositiva › escoge el diseño que mejor cuente tu idea. Para cambiar uno existente: Inicio › Diseño.'),
        ('Escribe en los marcadores', 'Reemplaza los textos de ejemplo. Tamaños, colores y espaciados ya están definidos: no los modifiques.'),
        ('Inserta tus imágenes', 'Haz clic en el ícono del marcador de imagen. La foto se recorta sola; ajústala con Recortar › Rellenar.'),
        ('Revisa y depura', 'Una idea por diapositiva. Si algo no cabe, divide la slide. Si te desordenas: Inicio › Restablecer.')]):
    fill(s, 50 + i, f'0{i + 1}'); fill(s, 54 + i, h); fill(s, 58 + i, b)

# --- Colores
s = new('26', 'Colores aprobados. Usa siempre los valores HEX exactos; están cargados como colores del tema.')
fill(s, 10, 'GUÍA RÁPIDA · COLOR'); fill(s, 'title', 'Colores aprobados')
main = [(PURPLE, 'Morado VERNE', 'Acento principal: cifras, destacados, eyebrows.', WHITE),
        (NAVY, 'Azul noche', 'Títulos y fondos oscuros.', WHITE),
        (LAV, 'Lavanda', 'Números de sección, acentos sobre oscuro.', NAVY),
        (TXT, 'Grafito', 'Texto de cuerpo.', WHITE),
        (WHITE, 'Blanco', 'Fondo principal y espacio negativo.', NAVY)]
for i, (c, name, use, tc) in enumerate(main):
    x = M + i * 2.43
    shape(s, x, 2.45, 2.25, 2.0, c, line=LINE if c == WHITE else None)
    tx(s, x + 0.2, 3.75, 1.9, 0.6, [[('#' + c, 13, tc, True)]], anchor='b')
    tx(s, x, 4.6, 2.25, 0.35, [[(name, 13, NAVY, True)]])
    tx(s, x, 4.95, 2.2, 0.7, [[(use, 10, GRAY)]])
eyebrow(s, M, 5.85, 6, 'Colores de apoyo')
for i, (c, name) in enumerate([(LAVBG, 'Lavanda fondo'), (LINE, 'Líneas'), (MID, 'Morado medio · gráficos'), (GRAY, 'Texto secundario'), (DEEP, 'Fondo profundo')]):
    x = M + i * 2.43
    shape(s, x, 6.2, 0.4, 0.4, c, MSO_SHAPE.OVAL)
    tx(s, x + 0.5, 6.18, 1.9, 0.45, [[('#' + c, 9.5, NAVY, True)], [(name, 8.5, GRAY)]], anchor='m')
tx(s, 8.0, 1.35, 4.68, 0.9, [[('Proporción sugerida: ', 10.5, NAVY, True), ('60% blanco · 25% azul noche · 10% morado · 5% acentos.', 10.5, GRAY)]], anchor='b')

# --- Tipografía
s = new('26', 'Sistema tipográfico. Los marcadores de cada diseño ya aplican estos estilos.')
fill(s, 10, 'GUÍA RÁPIDA · TIPOGRAFÍA'); fill(s, 'title', 'Poppins, con una jerarquía fija')
rows = [('Eyebrow / categoría', 'CATEGORÍA · TEMA', (10.5, PURPLE, True, False, 300), 'Bold 10,5 pt · mayúsculas · espaciado 3 pt'),
        ('Título principal', 'Título de portada', (30, NAVY, True), 'Bold 46 pt · interlineado 0,92'),
        ('Título de slide', 'Título de diapositiva', (24, NAVY, True), 'Bold 34 pt · interlineado 0,95'),
        ('Bajada', 'Bajada que da contexto en una o dos líneas.', (14, GRAY), 'Regular 16 pt · gris #6B6E80'),
        ('Cuerpo', 'Texto de desarrollo, en párrafos breves.', (12, TXT), 'Regular 12 pt · interlineado 1,25'),
        ('Destacado', 'La idea que debe recordarse.', (16, PURPLE, True), 'Bold 20 pt · morado'),
        ('Cifras', '+180%', (30, PURPLE, True), 'Bold 54–150 pt · morado'),
        ('Numeración', '01', (22, LAV, True), 'Bold 30 pt · lavanda'),
        ('Caption', 'Texto de apoyo o etiqueta.', (10, GRAY), 'Regular 10 pt'),
        ('Fuente / nota', 'Fuente: nombre de la fuente, año', (8.5, NOTE, False, True), 'Italic 8,5 pt · #9A9CAB')]
y = 2.3
for name, sample, f, spec in rows:
    hl(s, M, y, CW)
    tx(s, M, y + 0.06, 2.4, 0.38, [[(name, 10, GRAY, True)]], anchor='m')
    tx(s, 3.1, y + 0.02, 6.0, 0.42, [[(sample, *f)]], anchor='m')
    tx(s, 9.2, y + 0.06, 3.48, 0.38, [[(spec, 9, NAVY)]], anchor='m')
    y += 0.445

# --- Retícula y márgenes
s = new('26', 'Retícula. Las guías de edición (Vista › Guías) marcan los márgenes de 0,65", el inicio del eyebrow (1,0"), el fin del contenido (6,85") y el centro.')
fill(s, 10, 'GUÍA RÁPIDA · RETÍCULA'); fill(s, 'title', 'Márgenes y zonas')
shape(s, M, 2.35, CW, 4.5, PURPLE, alpha=5)
for x in (M + i * CW / 12 for i in range(1, 12)):
    shape(s, x, 2.35, 0.006, 4.5, LINE)
for x in (M, R): shape(s, x, 0, 0.012, 7.5, LAV)
for y in (1.0, 6.85): shape(s, 0, y, 13.333, 0.012, LAV)
tx(s, M + 0.15, 2.5, 6, 0.3, [[('ÁREA DE CONTENIDO · 12 COLUMNAS', 9, PURPLE, True, False, 200)]])
tx(s, M + 0.15, 6.55, 6, 0.25, [[('6,85" · fin del contenido. Debajo: numeración y web automáticas.', 8.5, GRAY)]])
tx(s, 7.2, 1.05, 4.6, 0.25, [[('1,0" · eyebrow  ·  1,28" · título', 8.5, PURPLE, True)]], align='r')
tx(s, 11.97, 0.82, 1.05, 0.2, [[('Logo fijo', 8, PURPLE, True)]], align='c')
tx(s, M + 4.3, 3.7, 7.4, 1.6, [[('0,65"', 40, PURPLE, True)], [('Margen lateral en todos los diseños. Las imágenes pueden ir a sangre; el texto, nunca fuera de los márgenes.', 12, NAVY)]])

# --- Logo
s = new('26', 'Uso del logo. El logo ya está en el Patrón de diapositivas: no lo dupliques ni lo muevas en cada slide.')
fill(s, 10, 'GUÍA RÁPIDA · LOGO'); fill(s, 'title', 'Uso del logo')
shape(s, M, 2.35, 3.85, 2.4, PANEL)
s.shapes.add_picture(IMG['logo'], I(M + 0.9), I(3.13), I(2.05), I(0.72))
shape(s, M + 4.09, 2.35, 3.85, 2.4, NAVY)
s.shapes.add_picture(IMG['logo_white'], I(M + 4.99), I(3.13), I(2.05), I(0.72))
cover_pic(s, PHOTO_AUDIENCE, M + 8.18, 2.35, 3.85, 2.4, 0.15)
shape(s, M + 8.18, 2.35, 3.85, 2.4, DEEP, alpha=45)
s.shapes.add_picture(IMG['logo_white'], I(M + 9.08), I(3.13), I(2.05), I(0.72))
for i, t in enumerate(['Morado sobre fondos claros', 'Blanco sobre fondos oscuros o morados', 'Blanco sobre foto, solo con zona oscura o overlay']):
    tx(s, M + i * 4.09, 4.85, 3.85, 0.35, [[(t, 11, NAVY, True)]])
hl(s, M, 5.45, CW)
tx(s, M, 5.6, 5.9, 1.1, [[('Sí', 11, PURPLE, True)], [('Respeta un espacio libre alrededor igual a la altura de la “v”. Úsalo en la esquina superior derecha (fijo en el patrón) o grande en portada y cierre.', 10.5, TXT)]])
tx(s, 6.84, 5.6, 5.84, 1.1, [[('No', 11, PURPLE, True)], [('No lo deformes, recolores, gires, ni le agregues sombras o contornos. No lo pongas sobre zonas de la foto con poco contraste.', 10.5, TXT)]])

# --- Tratamiento de imágenes
s = new('26', 'Seis formas de usar fotografía en el sistema VERNE. Cada miniatura es editable: copia la que necesites.')
fill(s, 10, 'GUÍA RÁPIDA · FOTOGRAFÍA'); fill(s, 'title', 'Tratamiento de imágenes')
tw, th_, gx, gy = 3.85, 1.85, 0.24, 0.62
demos = ['Imagen protagonista', 'Imagen lateral', 'Fotografía como fondo', 'Fotografía con overlay', 'Recorte de persona', 'Composición editorial']
for k, label in enumerate(demos):
    x = M + (k % 3) * (tw + gx); y = 2.2 + (k // 3) * (th_ + gy)
    if k == 0:
        cover_pic(s, PHOTO_MAN, x, y, tw, th_, 0.25)
    elif k == 1:
        shape(s, x, y, tw, th_, WHITE, line=LINE); cover_pic(s, PHOTO_WOMAN, x, y, tw * 0.45, th_, 0.3)
        tx(s, x + tw * 0.5, y + 0.45, tw * 0.45, 0.9, [[('Título', 13, NAVY, True)], [('Texto al costado de la foto', 8.5, GRAY)]])
    elif k == 2:
        cover_pic(s, PHOTO_AUDIENCE, x, y, tw, th_, 0.6)
        tx(s, x + 0.2, y + 1.05, 3, 0.6, [[('Mensaje sobre la foto', 13, WHITE, True)]], anchor='b')
    elif k == 3:
        cover_pic(s, PHOTO_AUDIENCE, x, y, tw, th_, 0.6)
        s.shapes.add_picture(A + 'fade_left.png', I(x), I(y), I(tw), I(th_))
        tx(s, x + 0.2, y + 0.55, 2.3, 0.9, [[('Overlay azul noche', 13, WHITE, True)], [('Garantiza legibilidad', 8.5, LINE)]])
    elif k == 4:
        shape(s, x, y, tw, th_, LAVBG)
        cover_pic(s, PHOTO_MAN, x + tw - 1.45, y, 1.25, th_, 0.2)
        tx(s, x + 0.2, y + 0.5, 2.0, 0.9, [[('Encuadre vertical', 13, NAVY, True)], [('La persona como protagonista', 8.5, GRAY)]])
    else:
        shape(s, x, y, tw, th_, PANEL)
        cover_pic(s, PHOTO_WOMAN, x + 0.2, y + 0.2, 1.5, th_ - 0.4, 0.3)
        cover_pic(s, PHOTO_AUDIENCE, x + 1.85, y + 0.2, 1.8, 0.85, 0.6)
        tx(s, x + 1.85, y + 1.15, 1.85, 0.5, [[('Retícula editorial', 11, NAVY, True)]])
    tx(s, x, y + th_ + 0.08, tw, 0.3, [[(label, 10.5, NAVY, True)]])
tx(s, M, 6.78, CW, 0.25, [[('Fotografía humana, real y contemporánea, con luz natural o tonos violeta sutiles. Evita stock evidente, gestos forzados y clichés de IA (robots, cerebros, redes de nodos).', 9, GRAY)]])

# --- Qué hacer / evitar
s = new('12', 'Diseño 12 Comparación.')
fill(s, 10, 'GUÍA RÁPIDA · CRITERIOS'); fill(s, 'title', 'Qué hacer y qué evitar')
fill(s, 40, 'HAZ'); fill(s, 41, 'Una idea clara por diapositiva')
fill(s, 42, [('Título + contexto mínimo + un elemento visual protagonista.', 0), ('Deja aire: el espacio negativo es parte de la marca.', 1),
             ('Usa los marcadores y estilos del patrón.', 1), ('Cifras grandes con su fuente.', 1), ('Fotos humanas con buen contraste para el texto.', 1)])
fill(s, 43, 'EVITA'); fill(s, 44, 'Slides que parecen documentos')
fill(s, 45, [('Párrafos largos: pásalos a las notas del presentador.', 0), ('Cajas y tarjetas para todo.', 1),
             ('Degradados, sombras, redes, nodos o íconos 3D.', 1), ('Cambiar fuentes, tamaños o colores fuera de la paleta.', 1), ('Texto sobre fotos sin overlay.', 1)])

# ======================================================================================
# SISTEMA VISUAL
# ======================================================================================
s = new('03', 'Diseño 03 Apertura de sección.')
fill(s, 12, '02'); fill(s, 'title', 'Sistema visual'); fill(s, 11, 'Recursos listos para copiar y pegar: todos son formas editables de PowerPoint.')

# --- Elementos gráficos
s = new('26', 'Kit de elementos. Copia y pega; mantén colores y grosores.')
fill(s, 10, 'SISTEMA VISUAL · ELEMENTOS'); fill(s, 'title', 'Elementos gráficos')
cols = [M, M + 3.07, M + 6.14, M + 9.21]
for x, t in zip(cols, ['Divisores y líneas', 'Etiquetas', 'Números de sección', 'Destacados']):
    eyebrow(s, x, 2.3, 2.9, t)
hl(s, cols[0], 2.85, 2.7); tx(s, cols[0], 2.95, 2.7, 0.25, [[('Línea fina · #D8D4EA · 0,75 pt', 8.5, GRAY)]])
hl(s, cols[0], 3.45, 2.7, PURPLE); tx(s, cols[0], 3.55, 2.7, 0.25, [[('Línea de acento · morado', 8.5, GRAY)]])
shape(s, cols[0], 4.05, 0.6, 0.03, LAV); tx(s, cols[0], 4.15, 2.7, 0.25, [[('Marca corta · portadas y cierre', 8.5, GRAY)]])
vl(s, cols[0] + 0.1, 4.6, 0.7); tx(s, cols[0] + 0.3, 4.75, 2.4, 0.4, [[('Separador vertical entre columnas', 8.5, GRAY)]])
chip(s, cols[1], 2.8, 1.5, 'Etiqueta', LAVBG, PURPLE)
chip(s, cols[1], 3.3, 1.5, 'Etiqueta', PURPLE, WHITE)
chip(s, cols[1], 3.8, 1.5, 'Etiqueta', WHITE, NAVY, line=LINE)
chip(s, cols[1], 4.3, 1.5, 'Etiqueta', NAVY, WHITE)
tx(s, cols[1], 4.8, 2.7, 0.5, [[('Para categorías, plazos o estados. Máximo 3 palabras.', 8.5, GRAY)]])
tx(s, cols[2], 2.75, 2.8, 0.7, [[('01', 40, LAV, True)]])
tx(s, cols[2], 3.5, 2.8, 0.55, [[('01  ', 14, PURPLE, True), ('Paso', 14, NAVY, True)]])
tx(s, cols[2], 4.05, 2.8, 0.6, [[('01', 30, PURPLE, True), (' / 03', 14, NOTE, True)]])
tx(s, cols[3], 2.8, 2.9, 0.9, [[('La idea que debe ', 14, NAVY, True), ('recordarse.', 14, PURPLE, True)]])
tx(s, cols[3], 3.75, 2.9, 0.7, [[('+180%', 34, PURPLE, True)]])
tx(s, cols[3], 4.45, 2.9, 0.4, [[('“Cita breve”', 14, NAVY, True)]])
hl(s, M, 5.55, CW)
eyebrow(s, M, 5.75, 3, 'Contenedores y formas')
shape(s, M, 6.1, 1.9, 0.7, PANEL); shape(s, M + 2.05, 6.1, 1.9, 0.7, LAVBG); shape(s, M + 4.1, 6.1, 1.9, 0.7, NAVY)
shape(s, M + 6.15, 6.1, 1.9, 0.7, None, line=LINE)
shape(s, M + 8.3, 6.1, 0.7, 0.7, PURPLE, MSO_SHAPE.OVAL); shape(s, M + 9.2, 6.1, 0.7, 0.7, None, MSO_SHAPE.OVAL, line=PURPLE)
shape(s, M + 10.1, 6.25, 1.9, 0.4, LAV, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
for x, t, c in [(M, 'Panel neutro', NAVY), (M + 2.05, 'Panel lavanda', NAVY), (M + 4.1, 'Panel oscuro', WHITE), (M + 6.15, 'Contorno', NAVY)]:
    tx(s, x + 0.15, 6.1, 1.7, 0.7, [[(t, 9, c, True)]], anchor='m')

# --- Fondos, overlays y marcos
s = new('26', 'Fondos, overlays y marcos de imagen. Los fondos también se pueden aplicar con Diseño › Formato del fondo.')
fill(s, 10, 'SISTEMA VISUAL · FONDOS Y MARCOS'); fill(s, 'title', 'Fondos, overlays y marcos')
eyebrow(s, M, 2.3, 6, 'Fondos')
bgs = [(WHITE, 'Blanco · base'), (LAVBG, 'Lavanda · statement'), (NAVY, 'Azul noche'), (PURPLE, 'Morado · secciones')]
for i, (c, t) in enumerate(bgs):
    x = M + i * 1.95
    shape(s, x, 2.65, 1.8, 1.1, c, line=LINE if c == WHITE else None)
    tx(s, x, 3.82, 1.8, 0.3, [[(t, 9, NAVY, True)]])
for i, (k, t) in enumerate([('glow_a', 'Brillo inferior'), ('glow_b', 'Brillo superior')]):
    x = M + (4 + i) * 1.95
    s.shapes.add_picture(IMG[k], I(x), I(2.65), I(1.8), I(1.1)); tx(s, x, 3.82, 1.8, 0.3, [[(t, 9, NAVY, True)]])
eyebrow(s, M, 4.4, 6, 'Overlays para texto sobre foto')
for i, (f, t) in enumerate([('fade_left.png', 'Lateral'), ('fade_bottom.png', 'Inferior')]):
    x = M + i * 1.95
    cover_pic(s, PHOTO_AUDIENCE, x, 4.75, 1.8, 1.4, 0.6)
    s.shapes.add_picture(A + f, I(x), I(4.75), I(1.8), I(1.4))
    tx(s, x, 6.22, 1.8, 0.3, [[(t, 9, NAVY, True)]])
x0 = M + 4.3
eyebrow(s, x0, 4.4, 6, 'Marcos de imagen')
frames = [(MSO_SHAPE.RECTANGLE, 'Rectangular', 2.0, 1.4), (MSO_SHAPE.ROUNDED_RECTANGLE, 'Redondeado', 2.0, 1.4),
          (MSO_SHAPE.OVAL, 'Circular · perfiles', 1.4, 1.4), (MSO_SHAPE.RECTANGLE, 'Vertical · persona', 0.95, 1.4)]
x = x0
for kind, t, w, h in frames:
    p = cover_pic(s, PHOTO_MAN if 'persona' in t or 'perfiles' in t else PHOTO_WOMAN, x, 4.75, w, h, 0.25)
    if kind != MSO_SHAPE.RECTANGLE:
        geom = 'roundRect' if kind == MSO_SHAPE.ROUNDED_RECTANGLE else 'ellipse'
        p._element.spPr.find(qn('a:prstGeom')).set('prst', geom)
    tx(s, x, 6.22, w + 0.3, 0.3, [[(t, 9, NAVY, True)]])
    x += w + 0.3

# --- Iconografía
s = new('26', 'Iconografía de línea (Lucide, licencia ISC). Más íconos: lucide.dev, exportados en #6D4ABE o blanco.')
fill(s, 10, 'SISTEMA VISUAL · ICONOGRAFÍA'); fill(s, 'title', 'Íconos de línea, simples y consistentes')
names = ['message-square', 'megaphone', 'users', 'user-round', 'search', 'chart-line', 'target', 'compass', 'lightbulb', 'layers',
         'graduation-cap', 'shield-check', 'globe', 'calendar', 'mail', 'phone', 'map-pin', 'file-text', 'pen-line', 'mic',
         'award', 'trending-up', 'book-open', 'briefcase']
for k, n in enumerate(names):
    x = M + (k % 12) * 0.86; y = 2.4 + (k // 12) * 0.95
    s.shapes.add_picture(f'{IC}{n}_p.png', I(x), I(y), I(0.5), I(0.5))
hl(s, M, 4.45, CW)
eyebrow(s, M, 4.65, 4, 'Variantes')
for i, (n, kind) in enumerate([('users', 'p'), ('target', 'n')]):
    s.shapes.add_picture(f'{IC}{n}_{kind}.png', I(M + i * 0.9), I(5.05), I(0.55), I(0.55))
for i, n in enumerate(['chart-line', 'shield-check']):
    s.shapes.add_picture(icon_disc(n), I(M + 1.9 + i * 1.0), I(5.0), I(0.65), I(0.65))
tx(s, M, 5.85, 4, 0.6, [[('Línea morada o azul noche sobre claro · ícono blanco en círculo morado sobre oscuro.', 9.5, GRAY)]])
tx(s, 6.84, 4.65, 5.84, 2.0, [[('Reglas', 11, PURPLE, True)],
                              [('Un solo estilo: línea de grosor uniforme, esquinas redondeadas.', 10.5, TXT)],
                              [('Tamaño mínimo 0,35". Un ícono por idea, nunca decorativo.', 10.5, TXT)],
                              [('Sin íconos 3D, con degradado, ni metáforas cliché de IA.', 10.5, TXT)]], after=4)

# --- Gráficos y tablas
s = new('26', 'Gráficos y tablas nativos. Los gráficos toman los colores del tema; la tabla usa el estilo “VERNE · Tabla editorial”, predeterminado en esta plantilla.')
fill(s, 10, 'SISTEMA VISUAL · DATOS'); fill(s, 'title', 'Gráficos y tablas')


def verne_chart(s, x, y, w, h, cats, vals, name='Serie', fmt='+0"%"'):
    cd = CategoryChartData(); cd.categories = cats; cd.add_series(name, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(x), I(y), I(w), I(h), cd)
    ch = gf.chart; ch.has_legend = False; ch.font.name = 'Poppins'; ch.font.size = Pt(10); ch.font.color.rgb = RGBColor.from_string(GRAY)
    pl = ch.plots[0]; pl.gap_width = 70; pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = fmt; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(13); dl.font.bold = True; dl.font.color.rgb = RGBColor.from_string(NAVY)
    ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGBColor.from_string(PURPLE)
    va = ch.value_axis; va.visible = False; va.has_major_gridlines = False
    ca = ch.category_axis; ca.format.line.color.rgb = RGBColor.from_string(LINE); ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(10); ca.tick_labels.font.color.rgb = RGBColor.from_string(GRAY)
    return gf


verne_chart(s, M, 2.35, 5.6, 3.6, ['Interacciones', 'Publicaciones propias', 'Seguidores'], [180, 256, 13], 'PetroTal')
tx(s, M, 6.05, 5.6, 0.5, [[('Ejemplo con datos reales: PetroTal · Embajadores. Una sola serie en morado; sin ejes ni cuadrícula innecesarios.', 9, GRAY)]])


def verne_table(s, x, y, w, rows, colw=None, h=None):
    gf = s.shapes.add_table(len(rows), len(rows[0]), I(x), I(y), I(w), I(h or 0.45 * len(rows)))
    tbl = gf.table
    gf._element.graphic.graphicData.tbl.tblPr.find(qn('a:tableStyleId')).text = T.TABLE_STYLE_ID
    tblPr = gf._element.graphic.graphicData.tbl.tblPr; tblPr.set('bandRow', '0'); tblPr.set('firstCol', '1')
    if colw:
        for i, cw in enumerate(colw): tbl.columns[i].width = I(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c); cell.text = val
            cell.margin_left = cell.margin_right = I(0.12); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                for rr in p.runs: rr.font.size = Pt(10.5 if r else 10); rr.font.name = 'Poppins'
    return gf


verne_table(s, 6.84, 2.35, 5.84, [['Frente', 'Consultoría', 'Formación'],
                                  ['Reputación y liderazgo', 'Marca Directiva y Corporativa', 'Talleres de Marca Directiva'],
                                  ['Visibilidad digital', 'Posicionamiento GEO', 'Formación para periodistas'],
                                  ['Adopción de IA', 'Diagnóstico de madurez digital', 'Taller IA Agéntica']], [1.9, 2.0, 1.94], 2.4)
tx(s, 6.84, 5.0, 5.84, 1.2, [[('Tabla: encabezado azul noche, líneas finas horizontales, primera columna en negrita. Sin bordes verticales ni rellenos alternos.', 9, GRAY)],
                             [('Paleta de gráficos (orden del tema): #6D4ABE · #C0B0F4 · #25283D · #8F72D6 · #D8D4EA · #6B6E80', 9, NAVY, True)]], after=6)

# ======================================================================================
# CATÁLOGO DE DISEÑOS (ejemplos con contenido real de VERNE)
# ======================================================================================
s = new('03', 'Diseño 03 Apertura de sección.')
fill(s, 12, '03'); fill(s, 'title', 'Catálogo de diseños')
fill(s, 11, 'Un ejemplo por diseño. Duplica el que necesites o créalo desde Nueva diapositiva.')

s = new('01', 'Diseño 01 Portada principal. Para portadas sin fotografía; el fondo con brillo ya está en el diseño.')
fill(s, 10, 'CREDENCIALES · 2026')
fill(s, 'title', 'Consultora Estratégica de Comunicación Digital, potenciada con Inteligencia Artificial')
fill(s, 1, 'Reputación y liderazgo · Visibilidad y presencia digital · Adopción y capacidades de IA'); fill(s, 11, 'www.verne.la')

s = new('02', 'Diseño 02 Portada con imagen. Usa fotos con el sujeto hacia la izquierda o centro.')
photo(s, 13, PHOTO_AUDIENCE)
fill(s, 10, 'FORMACIÓN · 2026'); fill(s, 'title', 'Habilitamos a tu equipo para que trabaje con IA')
fill(s, 1, 'Temario y casos a medida para cada empresa.'); fill(s, 11, 'Propuesta · Cliente · Fecha')

s = new('03', 'Diseño 03 Apertura de sección.')
fill(s, 12, '02'); fill(s, 'title', 'Visibilidad y presencia digital')
fill(s, 11, 'Logramos que te encuentren, citen y recomienden en buscadores, IA y medios.')

s = new('04', 'Diseño 04 Título + bajada. El diseño más flexible para contenido general.')
fill(s, 10, 'QUIÉNES SOMOS'); fill(s, 'title', 'El nexo entre la comunicación y la IA')
fill(s, 11, 'La comunicación estratégica ya no vive solo en los medios. Hoy se juega en tres frentes:')
fill(s, 14, [('Reputación y liderazgo', 1), ('Visibilidad y presencia digital', 1), ('Adopción y capacidades de IA', 1)])

s = new('05', 'Diseño 05 Texto + imagen. Imagen enmarcada a la derecha; el logo queda visible.')
photo(s, 13, PHOTO_MAN)
fill(s, 10, 'REPUTACIÓN Y LIDERAZGO'); fill(s, 'title', 'Marca Corporativa y Marca Directiva')
fill(s, 11, 'Construimos la marca de tu empresa a través de la voz de sus líderes.')
fill(s, 15, 'Los líderes cambian, la marca permanece.')
fill(s, 14, [('Diagnóstico · Estrategia · Gestión', 0)])

s = new('06', 'Diseño 06 Imagen + texto. Foto vertical a sangre a la izquierda.')
photo(s, 13, PHOTO_WOMAN)
fill(s, 10, 'REPUTACIÓN Y LIDERAZGO'); fill(s, 'title', 'Embajadores Digitales')
fill(s, 11, 'Amplificamos la comunicación corporativa a través de tus colaboradores.')
fill(s, 14, [('Estrategia: seleccionamos y activamos embajadores alineados a los objetivos de la marca.', 1),
             ('Acompañamiento: buenas prácticas, benchmark y co-creación de contenidos.', 1),
             ('Medición: informe mensual de desempeño, con Data Pulse.', 1)])

s = new('07', 'Diseño 07 Imagen a pantalla completa. El overlay es una imagen editable sobre la foto: mantenlo detrás del texto.')
pic = photo(s, 13, PHOTO_AUDIENCE)
ov = s.shapes.add_picture(A + 'fade_left.png', 0, 0, prs.slide_width, prs.slide_height)
pic._element.addnext(ov._element)
fill(s, 10, 'REPUTACIÓN Y LIDERAZGO'); fill(s, 'title', 'Los líderes cambian, la marca permanece.')
fill(s, 11, 'Construimos la marca de tu empresa a través de la voz de sus líderes.')

s = new('08', 'Diseño 08 Gran titular / statement. Una frase, sin más elementos.')
fill(s, 10, 'NUESTRA POSTURA'); fill(s, 'title', 'Ni solo tecnología ni solo comunicación tradicional.')
fill(s, 11, 'El nexo entre la comunicación y la IA.')

s = new('09', 'Diseño 09 Cifra protagonista. Siempre con fuente.')
fill(s, 10, 'QUIÉNES SOMOS · FRENTE 2'); fill(s, 'title', 'Visibilidad y presencia digital')
fill(s, 16, '70%'); fill(s, 17, 'de los consumidores peruanos usa IA')
fill(s, 14, 'para informarse antes de comprar o contratar.'); fill(s, 18, 'Fuente: Impronta Research, ago 2025')

s = new('10', 'Diseño 10 Dos columnas.')
fill(s, 10, 'REPUTACIÓN Y LIDERAZGO'); fill(s, 'title', 'Marca Corporativa y Marca Directiva')
fill(s, 11, 'Construimos la marca de tu empresa a través de la voz de sus líderes.')
fill(s, 20, 'Marca Corporativa'); fill(s, 22, 'La reputación y credibilidad de la empresa en el ecosistema digital. Es lo que permanece.')
fill(s, 21, 'Marca Directiva'); fill(s, 23, 'Tus C-level, directivos y voceros dan rostro y voz a esa narrativa. Es lo que la amplifica.')

s = new('11', 'Diseño 11 Tres columnas.')
fill(s, 10, 'EMBAJADORES DIGITALES'); fill(s, 'title', 'Cómo trabajamos')
fill(s, 11, 'Amplificamos la comunicación corporativa a través de tus colaboradores.')
for i, (h, b) in enumerate([('Estrategia', 'Seleccionamos y activamos embajadores alineados a los objetivos de la marca.'),
                            ('Acompañamiento', 'Buenas prácticas, benchmark y co-creación de contenidos, con soporte continuo.'),
                            ('Medición', 'Data pública, matriz de evaluación e informe mensual de desempeño, con Data Pulse.')]):
    fill(s, 30 + i, f'0{i + 1}'); fill(s, 33 + i, h); fill(s, 36 + i, b)

s = new('12', 'Diseño 12 Comparación.')
fill(s, 10, 'QUÉ HACEMOS'); fill(s, 'title', 'Nuestros servicios')
fill(s, 40, 'CONSULTORÍA'); fill(s, 41, 'Lo hacemos por ti')
fill(s, 42, [('Marca Directiva y Corporativa', 1), ('Posicionamiento GEO', 1), ('Diagnóstico de madurez digital', 1), ('Agentic AI Studio · GEO Rank · Data Pulse', 1)])
fill(s, 43, 'FORMACIÓN'); fill(s, 44, 'Habilitamos a tu equipo')
fill(s, 45, [('Talleres de Marca Directiva y Embajadores', 1), ('Formación para periodistas', 1), ('Talleres de IA, data y automatización', 1), ('Claude AI para empresas · Taller IA Agéntica', 1)])

s = new('13', 'Diseño 13 Proceso / pasos. Si usas tres pasos, elimina los marcadores del cuarto.')
fill(s, 10, 'REPUTACIÓN Y LIDERAZGO'); fill(s, 'title', 'Marca Corporativa y Directiva: el método')
fill(s, 11, 'Para que el mensaje sea de la empresa y no dependa de una sola persona.')
for i, (h, b) in enumerate([('Diagnóstico', 'Analizamos la presencia digital de la empresa y de sus directivos, y la comparamos con la de sus pares.'),
                            ('Estrategia', 'Definimos la narrativa corporativa, los ejes de contenido y el rol de cada vocero.'),
                            ('Gestión', 'Creamos contenidos, acompañamos a los voceros y medimos el desempeño.')]):
    fill(s, 50 + i, f'0{i + 1}'); fill(s, 54 + i, h); fill(s, 58 + i, b)
drop(s, 53, 57, 61)

s = new('14', 'Diseño 14 Timeline. Ejemplo con el plan GEO (15 + 15 + 60 días).')
fill(s, 10, 'POSICIONAMIENTO GEO'); fill(s, 'title', 'Un plan de 90 días')
fill(s, 11, 'Logramos que la IA cite y recomiende a tu marca.')
for i, (d, h, b) in enumerate([('Días 1–15', 'Diagnóstico', 'Cómo aparece tu marca frente a la competencia, con brechas y oportunidades.'),
                               ('Días 16–30', 'Estrategia', 'Narrativas, arquitectura de contenidos y recorrido del usuario en la IA.'),
                               ('Días 31–90', 'Implementación', 'Parrilla GEO, producción de contenidos y verificación de citación.'),
                               ('Día 90', 'Resultados', 'Medibles en visibilidad, posición, citaciones y share of voice en IA.')]):
    fill(s, 70 + i, d); fill(s, 75 + i, h); fill(s, 80 + i, b)

s = new('15', 'Diseño 15 Diagrama / ecosistema. Anillos concéntricos para mostrar niveles o capas.')
fill(s, 10, 'REPUTACIÓN Y LIDERAZGO'); fill(s, 'title', 'La marca, del centro hacia afuera')
fill(s, 11, 'Cada capa amplifica a la anterior.')
fill(s, 14, [('Marca Corporativa: lo que permanece.', 1), ('Marca Directiva: lo que la amplifica.', 1), ('Embajadores: la voz más genuina de la marca.', 1)])
fill(s, 90, 'Marca Corporativa'); fill(s, 91, 'Marca Directiva'); fill(s, 92, 'Embajadores digitales')

s = new('16', 'Diseño 16 Servicios o soluciones. Inserta en cada círculo un ícono de la guía (círculo morado con ícono blanco).')
fill(s, 10, 'ADOPCIÓN Y CAPACIDADES DE IA'); fill(s, 'title', 'Nuestras cuatro soluciones de Inteligencia Artificial')
sols = [('pen-line', 'Agentic AI Studio', 'Estudio de contenidos con agentes de IA', 'Desde el ADN de tu marca, los agentes crean piezas multiformato listas para aprobar en minutos.'),
        ('shield-check', 'Agentic AI Reputation', 'Consultor de reputación', 'Asesora ante cualquier situación que afecte tu reputación.'),
        ('search', 'GEO Rank', 'Plataforma de monitoreo GEO', 'Mide la visibilidad, las citas y la veracidad de tu marca en las IAs.'),
        ('chart-column', 'Data Pulse', 'Tracking de embajadores digitales', 'Mide el desempeño de tus embajadores y convierte la data en acciones.')]
for i, (ic, n, t, d) in enumerate(sols):
    photo(s, 100 + i, icon_disc(ic)); fill(s, 104 + i, n); fill(s, 108 + i, t); fill(s, 112 + i, d)

s = new('17', 'Diseño 17 Casos / proyectos.')
photo(s, 13, PHOTO_WOMAN)
fill(s, 10, 'CASO · REPUTACIÓN Y LIDERAZGO'); fill(s, 'title', 'PetroTal · Embajadores')
fill(s, 11, 'Amplificamos la comunicación corporativa a través de sus colaboradores.')
fill(s, 14, [('Estrategia, acompañamiento y medición con Data Pulse.', 0)])
for i, (n, l) in enumerate([('+180%', 'interacciones'), ('+256%', 'publicaciones propias'), ('+13%', 'seguidores (+6,365)')]):
    fill(s, 120 + i, n); fill(s, 123 + i, l)

s = new('18', 'Diseño 18 Clientes / logos. Inserta logos en PNG transparente, en gris o en su color original, de forma consistente.')
fill(s, 10, 'CON QUIÉNES TRABAJAMOS'); fill(s, 'title', 'Confían en nosotros')
fill(s, 11, 'Empresas líderes de sectores regulados y competitivos.')
groups = [('Energía, Minería e Hidrocarburos', ['image3', 'image6', 'image7', 'image9', 'image11']),
          ('Banca y Servicios Financieros', ['image16', 'image17', 'image18', 'image19', 'image20']),
          ('Consumo, Retail e Industria', ['image21', 'image22', 'image23', 'image24'])]
for r, (name, logos) in enumerate(groups):
    fill(s, 130 + r, name)
    for c, lg in enumerate(logos):
        ph = s.placeholders[140 + r * 5 + c]
        im = Image.open(MEDIA + lg + '.png'); ar = im.width / im.height
        bw, bh = ph.width, ph.height
        w = min(bw, int(bh * 0.6 * ar)); h = int(w / ar)
        x, y = ph.left + (bw - w) // 2, ph.top + (bh - h) // 2
        p_ = s.shapes.add_picture(MEDIA + lg + '.png', x, y, w, h)
        ph._element.addprevious(p_._element); ph._element.getparent().remove(ph._element)
    for c in range(len(logos), 5): drop(s, 140 + r * 5 + c)

s = new('19', 'Diseño 19 Equipo / perfiles. Para dos perfiles, elimina los marcadores del tercero.')
fill(s, 10, 'EQUIPO'); fill(s, 'title', 'Nosotros te acompañamos')
fill(s, 11, 'Potenciamos la comunicación digital en empresas, mejorando el posicionamiento y reputación de directivos y equipos.')
for i, (img, n, rl, b) in enumerate([('image51', 'Víctor Lozano', 'DIRECTOR DE INNOVACIÓN', '+20 años de experiencia en Investigación de Mercados, Marketing Digital, Gestión de Marcas Directivas y Reputación.'),
                                     ('image55', 'Julio Talledo', 'DIRECTOR COMERCIAL & MARKETING', 'Especialista en Marketing e Inteligencia Artificial Generativa. Docente de IA Generativa en UDEP, CENTRUM PUCP, UPC y USIL.')]):
    photo(s, 160 + i, MEDIA + img + '.png'); fill(s, 163 + i, n); fill(s, 166 + i, rl); fill(s, 169 + i, b)
drop(s, 162, 165, 168, 171)

s = new('20', 'Diseño 20 Testimonio o cita. Usa solo citas reales y autorizadas, con nombre y cargo.')
fill(s, 170, 'Los líderes cambian, la marca permanece.')
fill(s, 172, 'VERNE'); fill(s, 173, 'Comunicación e Innovación')
drop(s, 171)
for i, (y, h) in ((172, (5.5, 0.4)), (173, (5.92, 0.3))):
    p_ = s.placeholders[i]; p_.left, p_.top, p_.width, p_.height = I(M), I(y), I(7), I(h)

s = new('21', 'Diseño 21 Datos / métricas. Hasta cuatro cifras con su contexto.')
fill(s, 10, 'CASO · MARCA DIRECTIVA'); fill(s, 'title', 'Enel Perú · Marca Directiva')
fill(s, 11, 'Resultados medidos con nuestros clientes.')
for i, (n, l) in enumerate([('+2 pts', 'Posicionamiento digital'), ('+68%', 'Seguidores'), ('+35%', 'Contenido propio'), ('+27%', 'Interacciones')]):
    fill(s, 180 + i, n); fill(s, 184 + i, l)
drop(s, 188, 189, 190, 191, 192)

s = new('22', 'Diseño 22 Tabla simple. Inserta la tabla desde el ícono del marcador; toma el estilo VERNE automáticamente.')
fill(s, 10, 'QUÉ HACEMOS'); fill(s, 'title', 'Nuestros servicios')
fill(s, 11, 'Consultoría (lo hacemos por ti) y Formación (habilitamos a tu equipo) en cada frente.')
gf = s.placeholders[200].insert_table(4, 3)
gf._element.graphic.graphicData.tbl.tblPr.find(qn('a:tableStyleId')).text = T.TABLE_STYLE_ID
tp = gf._element.graphic.graphicData.tbl.tblPr; tp.set('bandRow', '0'); tp.set('firstCol', '1')
data = [['Frente', 'Consultoría', 'Formación'],
        ['Reputación y liderazgo', 'Marca Directiva y Corporativa · Embajadores digitales', 'Talleres de Marca Directiva y Embajadores'],
        ['Visibilidad y presencia digital', 'Posicionamiento GEO · Marketing de Contenidos · Asuntos Públicos', 'Formación para periodistas · Political Intelligence AI'],
        ['Adopción y capacidades de IA', 'Diagnóstico de madurez digital · Agentic AI Studio · GEO Rank · Data Pulse', 'Talleres de IA, data y automatización · Claude AI para empresas']]
for r, row in enumerate(data):
    for c, v in enumerate(row):
        cell = gf.table.cell(r, c); cell.text = v; cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = I(0.12)
        for p in cell.text_frame.paragraphs:
            for rr in p.runs: rr.font.size = Pt(11); rr.font.name = 'Poppins'
for i, w in enumerate((3.0, 4.6, 4.43)): gf.table.columns[i].width = I(w)
drop(s, 201)

s = new('23', 'Diseño 23 Gráfico. Inserta el gráfico desde el ícono del marcador: toma los colores del tema.')
fill(s, 10, 'CASO · EMBAJADORES'); fill(s, 'title', 'PetroTal · Embajadores')
fill(s, 11, 'Resultados medidos con nuestros clientes.')
cd = CategoryChartData(); cd.categories = ['Interacciones', 'Publicaciones propias', 'Seguidores']; cd.add_series('PetroTal', [180, 256, 13])
gfc = s.placeholders[210].insert_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
ch = gfc.chart; ch.has_legend = False; ch.font.name = 'Poppins'; ch.font.size = Pt(10)
pl = ch.plots[0]; pl.gap_width = 70; pl.has_data_labels = True
pl.data_labels.number_format = '+0"%"'; pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.data_labels.font.size = Pt(13); pl.data_labels.font.bold = True
pl.data_labels.font.color.rgb = RGBColor.from_string(NAVY)
pl.series[0].format.fill.solid(); pl.series[0].format.fill.fore_color.rgb = RGBColor.from_string(PURPLE)
ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
ch.category_axis.format.line.color.rgb = RGBColor.from_string(LINE)
ch.category_axis.tick_labels.font.color.rgb = RGBColor.from_string(GRAY)
fill(s, 15, '+256% en publicaciones propias'); fill(s, 14, 'La cifra con mayor crecimiento del caso, junto a +180% en interacciones y +13% en seguidores (+6,365).')
drop(s, 211)

s = new('24', 'Diseño 24 Cierre.')
fill(s, 'title', 'Gracias'); fill(s, 11, 'Consultora Estratégica de Comunicación Digital, potenciada con Inteligencia Artificial.')

s = new('25', 'Diseño 25 Contacto.')
photo(s, 13, PHOTO_AUDIENCE)
fill(s, 10, 'CONTACTO'); fill(s, 'title', 'Conversemos')
fill(s, 11, 'Nosotros te acompañamos.')
for i, (l, v) in enumerate([('WEB', 'www.verne.la'), ('VÍCTOR LOZANO · DIRECTOR DE INNOVACIÓN', 'linkedin.com/in/victorlozanou'),
                            ('JULIO TALLEDO · DIRECTOR COMERCIAL & MARKETING', 'linkedin.com/in/juliotalledo')]):
    fill(s, 220 + i, l); fill(s, 223 + i, v)

s = new('26', 'Diseño 26 Solo título: para composiciones libres respetando la retícula.')
fill(s, 10, 'DISEÑO LIBRE'); fill(s, 'title', 'Solo título')
tx(s, M, 2.4, 7, 1.2, [[('Usa este diseño cuando ninguno de los anteriores se ajuste. Mantén los márgenes de 0,65", los estilos tipográficos y la paleta.', 14, GRAY)]])

# ---- transitions: subtle fade everywhere
for sl in prs.slides:
    el = sl._element
    tr = el.makeelement(qn('p:transition'), {'spd': 'med'}); tr.append(tr.makeelement(qn('p:fade'), {}))
    anchor = el.find(qn('p:clrMapOvr'))
    (anchor if anchor is not None else el.find(qn('p:cSld'))).addnext(tr)

prs.core_properties.title = 'Plantilla maestra VERNE'
prs.core_properties.author = 'VERNE · Comunicación e Innovación'
prs.save(OUT + '.pptx')

# ---- .potx: same package, template content type
with zipfile.ZipFile(OUT + '.pptx') as zi, zipfile.ZipFile(OUT + '.potx', 'w', zipfile.ZIP_DEFLATED) as zo:
    for item in zi.infolist():
        data = zi.read(item.filename)
        if item.filename == '[Content_Types].xml':
            data = data.replace(b'presentationml.presentation.main+xml', b'presentationml.template.main+xml')
        zo.writestr(item, data)
print('slides', len(prs.slides), 'layouts', len(prs.slide_layouts))
